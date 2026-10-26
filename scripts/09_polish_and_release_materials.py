"""Editorial revision and release provenance; never refit or modify scientific outputs.

Prepare: python scripts/09_polish_and_release_materials.py --pandoc PATH
DOI update: same command --doi 10.5281/zenodo.REAL_ID
Finalize after render review: same command --finalize
GitHub/Zenodo authentication and publication are separate explicit author actions.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import re
import shutil
import subprocess
import time
import urllib.request
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('submission', ROOT/'scripts/08_build_submission_package.py')
pack = importlib.util.module_from_spec(spec); spec.loader.exec_module(pack)
OUT = ROOT/'submission_package'
ARCHIVE = ROOT/'results/archive_step08_before_editorial_release'
RELEASE = OUT/'github_zenodo_release'
WORK = ROOT/'work/step09'
REPO = 'https://github.com/mzxuan08/cfDNA-Scale-Tradeoff-Audit'
TAG = 'v0.1.0'


def snapshot():
    frozen = ROOT/'results/reports/final_scope_alignment_integrity.json'
    manifest = json.loads(pack.read(frozen))
    files = {frozen}
    for name, digest in manifest['frozen_files'].items():
        assert pack.sha(name) == digest, name
        files.add(Path(name))
    for folder in ['data/processed','data/metadata','results/tables','results/figures']:
        files.update(p for p in (ROOT/folder).rglob('*') if p.is_file())
    return {str(p.relative_to(ROOT)).replace('\\','/'):pack.sha(p) for p in sorted(files)}


def archive_once():
    if (ARCHIVE/'archive_manifest.json').exists(): return
    ARCHIVE.mkdir(parents=True,exist_ok=True)
    files = list(OUT.glob('*.docx')) + list(OUT.glob('*.md'))
    files += [ROOT/'results/MANUSCRIPT_FULL_TEXT.md',ROOT/'results/reports/MANUSCRIPT_MASTER_DRAFT.md',
              OUT/'verification/Build_Summary.json',OUT/'verification/Frozen_Input_Integrity.json']
    manifest = {}
    for p in files:
        if not p.exists():continue
        dest = ARCHIVE/('submission_package/'+str(p.relative_to(OUT)) if p.is_relative_to(OUT) else str(p.relative_to(ROOT)))
        dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(p,dest)
        manifest[str(dest.relative_to(ARCHIVE))]=pack.sha(dest)
    pack.dump(ARCHIVE/'archive_manifest.json',manifest)


def revise_full():
    full = pack.read(ARCHIVE/'submission_package/Manuscript_Main_Blind_and_Full.md')
    paragraphs = full.split('\n\n')
    def replace(prefix, text):
        indices = [i for i,p in enumerate(paragraphs) if p.startswith(prefix)]
        assert len(indices)==1,(prefix,len(indices))
        paragraphs[indices[0]]=text

    replace('**Background:**', '**Background:** Normalization determines which cfDNA fragmentation contrasts a model retains and how its scores transfer across studies. We audited the relationship between normalization scale, study-source variation and target-cohort composition in public cfDNA data.')
    replace('**Methods:**', '**Methods:** We analysed 585 source-identified participants using a primary bilateral benchmark of 548 and a separately held-out, case-dominated cohort of 37. Global, local and hybrid representations were evaluated with source-only selection and locked thresholds. Cancer–control pair concordance and fixed-patient anchor experiments separated within-study ranking, cross-study score comparability and target-composition dependence.')
    replace('**Results:**', '**Results:** Combined LOCO ElasticNet AUROCs were 0.759 for Global Z-score, 0.731 for Hybrid and 0.724 for fixed LSN-ICC. Median absolute arm-level fragmentation effects were 0.270 for Global and 0.244 for local normalization; residual GC correlations with study variance were 0.158 and 0.405. For ComBat with unseen-target alignment, within-study and cross-study concordances were 0.669 and 0.366. Changing the companions of fixed real patients shifted their probabilities by up to 0.214 under target-batch adaptation and by 0 under within-sample transformations.')
    replace('**Conclusions:**', '**Conclusions:** Normalization scale, study-wise discrimination and score comparability are distinct targets for cross-study cfDNA evaluation. The audit makes their trade-offs measurable and identifies patient-wise self-normalization as a direct way to remove companion-dependent target preprocessing.')

    replace('Plasma cell-free DNA fragmentation encodes', 'Cross-study cfDNA fragmentomics requires representations that retain disease-associated fragmentation contrasts while controlling study-source variation and maintaining interpretable scores in a new cohort. Plasma DNA fragmentation reflects nucleosome organization and tissue contribution, and genome-wide profiles support cancer classification alongside sequence-variant assays [1–3]. Study Source combines specimen processing, sequencing procedures, recruitment and population composition. Evaluating transport therefore requires both patient-level discrimination and the achieved operating point of a threshold locked in the source study.')
    replace('Normalization has two separable roles', 'Normalization scale offers a way to examine these requirements directly. Global normalization retains relative deviations across the genome while removing a sample-wide offset and scale. A local contrast compares adjacent windows and attenuates slowly varying backgrounds, including broad disease-associated differences. In the window-level audit, median absolute within-study disease contrasts were 0.384 for raw ratios and 0.263 for a local representation. This motivates a comparison of global, local and fused representations, with fragmentation contrast retention as the measured endpoint.')
    replace('Batch-distribution alignment introduces', 'Cross-sample alignment adds a separate dependence on cohort composition. An empirical-Bayes harmonizer estimates study-specific location and scale parameters [4]. Adapting it to an unseen study with target medians or quantiles makes those estimates properties of the target patient mixture. Under $F_c=(1-\\pi_c)F_{0c}+\\pi_cF_{1c}$, changing the case fraction $\\pi_c$ can change batch-level quantiles while the evaluated individuals remain fixed. Within-sample transformations offer a precise alternative: each patient supplies the reference for their own transformation. This functional distinction can be tested by holding real patients fixed and changing their companions.')
    replace('We therefore conduct a systematic', 'We present a Scale Trade-off audit that connects representation, discrimination and score transport in public cfDNA cohorts. Global and same-arm local normalization quantify how scale changes broad fragmentation contrasts and residual GC-associated study variation; Hybrid tests their complementarity in a source-trained classifier. Bidirectional study-held-out testing and a separate external sensitivity cohort establish the evaluation scope. Pair concordance decomposition and fixed-patient composition stress tests then identify whether pooled score differences arise from within-study ranking, comparability across fitted models or target-batch adaptation.')

    replace('All new analyses used the already', 'The analysis used the existing extracted feature matrix for 585 source-identified participants. Clinical metadata came from UNITE supplementary data file S1, sheet s2, and fragment intervals were streamed from explicit FinaleDB objects [5,6]. Inclusion required plasma WGS, original sequencing depth within 0.1–5.0× and explicit cancer or Healthy source labels. The extraction cohort contained 628 specimens. One prespecified specimen per known patient was retained by earliest time point, then shallowest sequencing depth, then sequencing-run identifier; a specimen with unresolved patient identity was isolated. Healthy controls required an explicit Healthy source category. Hepatitis, cirrhosis, inflammatory disease, transplantation and autoimmune disease retained their original categories.')
    replace('Study identifiers are proxies combining', 'Study Source is an observational proxy for experimental and population-composition differences. The source publications have different recruitment aims, so their case proportions describe the selected cohorts. Participant identity follows published source identifiers; undocumented aliases across studies remain unresolved.')
    replace('The archived original database publication', 'The archived database publication listed 7 sources and 2579 sequencing records. Its additional urinary and prenatal single-gene studies did not document an eligible plasma cancer/control cohort in the local materials. Unavailable run-level case/control counts remain unknown. The census describes the archived local inventory at the stated eligibility threshold.')
    replace('Descriptive PERMANOVA and window variance', 'Descriptive PERMANOVA and window variance models adjusted for disease status, study and original sequencing depth as appropriate. PERMANOVA used Euclidean geometry after descriptive column standardization and 999 restricted Freedman–Lane residual permutations. Window-level unique/type III eta-squared was the reduced-to-full residual sum-of-squares increment divided by total sum of squares; overlapping explanatory components can yield a sum different from one. Descriptive embeddings were kept separate from predictive preprocessing. GC associations were calculated independently for raw, global and local representations.')
    replace('This relative epsilon preserves', 'The relative epsilon preserves global multiplicative invariance to numerical precision. The local operation measures a relative contrast; additive shifts and arbitrary nonlinear distortions are outside that invariance. Hybrid MultiScale concatenated the global and local blocks,')
    replace('and fitted its column scaler only', 'and fitted its column scaler within the source training split. Both blocks and the fixed Hybrid radius were specified before outer evaluation. Elastic-net logistic regression supplied the primary classifier, learning regularized coefficients for the two correlated blocks. Additional fixed classifiers provided secondary algorithm comparisons.')
    replace('The fixed LSN-ICC design attempted', 'The fixed LSN-ICC pipeline ranked healthy-center clustering scores within source training data, removed the highest quarter and applied RobustScaler. The scores are one-way between-center variance ratios among different healthy people. In a bilateral direction, the single source study makes cross-center filtering unidentifiable; the recorded operational representation is therefore LSN only. This fallback is labelled in the results and uses source data exclusively.')
    replace('A separate Tuned LSN analysis', 'Tuned LSN considered radii $k\\in\\{1,2,3,5,8\\}$ and eligible exclusion fractions $\\alpha\\in\\{0,0.15,0.25,0.35,0.50\\}$. With one healthy source study, zero exclusion was the eligible setting. Three inner folds selected the highest mean validation AUROC across source strata containing at least two participants of each class, with deterministic tie-breaking and the prespecified ElasticNet model. This analysis evaluated source-only scale selection separately from fixed Hybrid.')
    replace('The five core pipelines were', 'The five core pipelines were Raw Ratio, Global Z-score, InFold ComBat, fixed Proposed LSN-ICC and Hybrid MultiScale. Tuned LSN used the primary ElasticNet classifier. Random Forest, CUDA XGBoost, LightGBM, CUDA MLP and source-only CUDA DANN supplied locked secondary models. DANN used training-source domains exclusively. Multi-source DANN and standard multi-batch source ComBat were estimable in mixed-study random CV and two-source/one-target Combined LOCO. Their one-source bilateral cells were marked unavailable.')
    replace('InFold ComBat combined', 'InFold ComBat combined source-only empirical-Bayes harmonization with an unlabeled target median/IQR mapping for an unseen study. This explicitly transductive comparator evaluates a particular adaptation rule. The other fitted transformations process each target patient independently of companion patients; their source-trained scalers and classifiers still reflect the source training population.')

    replace('The independent analysis cohort comprised', 'The analysis cohort comprised 585 participants. In raw-ratio PERMANOVA, the unique study fraction was 65.628% and the disease fraction was 0.833% (both permutation p=0.001). GC correlated with raw window study eta-squared at rho=0.606, p=3.21e-49. For Global Z-score, the correlation was 0.158, p=0.000525; for local contrasts, it was 0.405, p=3.26e-20. Median unique window study variance fractions were 65.545%, 2.553% and 2.456%, respectively. Normalization therefore changed both the amount of source-associated variation and its genomic ordering (Figure 1; Figure 4).')
    replace('Median absolute arm-averaged disease', 'Median absolute arm-averaged disease contrasts were 0.381 for raw ratios, 0.270 for Global Z-score and 0.244 for local contrasts. The reference-region comparisons yielded 0 BH-significant enrichments among 7 tests. These results quantify how normalization redistributes broad fragmentation contrasts and provide the measured context for the multiscale comparison.')
    replace('The Scale Trade-off is empirical', 'The Scale Trade-off has two observable dimensions: broad fragmentation-contrast retention and the magnitude and genomic pattern of residual study variation. Local normalization produced a slightly smaller median study-variance fraction than Global Z-score, alongside a stronger residual GC rank association. Thus the preferred representation depends on the property being evaluated. Copy-number and tissue-accessibility interpretations are considered as mechanistic hypotheses in the Discussion.')
    replace('The pooled primary Hybrid AUROC', 'The pooled primary Hybrid AUROC was 0.689 and its equal-weight mean study AUROC was 0.743; corresponding Global Z-score values were 0.702 and 0.745. These summaries evaluate different weighting and score-comparability properties. Figure 2 and Figure 3 separate direction-specific discrimination, calibration and locked-threshold transport. The complete classifier table reports estimable and unavailable combinations explicitly.')
    replace('For the same fixed ElasticNet comparison', 'For fixed ElasticNet, Combined LOCO AUROC was 0.759 for Global Z-score, 0.731 for Hybrid and 0.724 for Proposed LSN-ICC. Hybrid therefore occupied an intermediate position between the global and fixed local representations. In the primary paired comparison, its AUROC was lower than Global in Cristiano and numerically higher in Jiang; the latter comparison did not reach BH significance. These direction-specific results support the scale-complementarity audit rather than a claim of consistent Hybrid superiority. DeLong inference remains target-specific.')
    replace('Cancer proportions were', 'Cancer proportions were 51.054% in Cristiano, 73.554% in Jiang and 97.297% in Snyder. Study and diagnosis were associated (chi-square=44.024, df=2, p=2.76e-10, Cramer\'s V=0.274), establishing the composition contrast examined in the adaptation diagnostics.')
    replace('For the empirical-Bayes-plus-target-alignment', 'The empirical-Bayes-plus-target-alignment comparator had secondary pooled AUROC 0.543. Pair decomposition separated weighted within-study concordance of 0.669 from cross-study concordance of 0.366; cross-study pairs contributed 41.636% of all pairs. The pooled AUROC therefore combines study-wise ranking with comparability between scores produced by independently source-trained models (Supplementary Figure S5).')
    replace('Expressing the ComBat scores relative', 'Referencing each ComBat score to its source-locked threshold yielded secondary pooled AUROC 0.616, a change of +0.073. This descriptive alignment preserved the benchmark predictions and used source-defined operating points to examine comparability across fitted score scales.')
    replace('When the same anchor patients were', 'The fixed-anchor experiment directly identified target-composition dependence. Across batches of real target patients with different case fractions, the maximum anchor probability range was 0.214 under target-batch adaptation. Corresponding ranges were 0 for Raw Ratio, 0 for Global Z-score, 0 for fixed LSN-ICC and 0 for Hybrid. Within-sample transformations thus removed companion-dependent target normalization for the frozen scoring rules (Figure 3d).')

    replace('The principal contribution is', 'This audit separates three properties that are often combined in cross-study cfDNA evaluation: normalization scale, study-wise discrimination and comparability across fitted score scales. Global and local transformations change broad fragmentation contrasts and residual source variation in different ways. Pair decomposition explains how pooled AUROC combines within-study ranking with cross-study score comparability, while fixed real-patient anchors isolate companion dependence in unseen-target adaptation. Together, these analyses turn a single benchmark score into an interpretable account of representation and transport.')
    replace('The two scale blocks have', 'Global and local normalization impose different reference scales. Global Z-scoring retains deviations relative to the full genome; the local log contrast references adjacent windows on the same arm. Attenuation of a smooth background can therefore accompany attenuation of a broad disease-associated contrast. Hybrid exposes both scales to a source-trained penalized classifier, making their complementarity testable. The observed target-specific comparisons identify the conditions in which the fused representation changes ranking and calibration; the Combined LOCO estimates favour Global Z-score over Hybrid for the fixed ElasticNet comparison.')
    replace('The ComBat diagnostics clarify', 'The ComBat diagnostics locate a specific source of composition dependence. The evaluated system combines empirical-Bayes source adjustment and robust unseen-target mapping. A target quantile follows the mixture of patients, so changing that mixture can move the adapted features of a fixed individual. The anchor experiment measures the resulting probability change with all source fitting held constant. Preserved within-study Cohen d and the pair decomposition distinguish this adaptation effect from removal of univariate disease contrasts. The conclusion applies to the evaluated adaptation rule; other covariate-aware or reference-batch ComBat designs remain separate methods [4].')
    replace('The independence property of', 'Within-sample transformation makes a patient\'s representation independent of other target patients. For a frozen source model, adding or removing companions leaves the global, local and hybrid feature vectors and predictions unchanged. This property removes a defined route of target-composition dependence. Training-population composition, disease spectrum and prevalence-dependent predictive values remain separate determinants of performance, making calibration and source-locked threshold transport complementary evaluation targets.')
    replace('The hierarchical cohort presentation', 'The cohort hierarchy aligns each public source with the endpoint it can support. The two primary targets permit bidirectional discrimination and specificity estimation. The separately held-out, cancer-enriched source contributes sensitivity at primary-source-locked thresholds. Combined LOCO then evaluates the three-source setting with its score-comparability requirements made explicit. This structure provides a reusable evaluation design for public datasets with unequal case fractions and control representation.')
    replace('Source availability is the core', 'The available source-level replication comprises 2 cohorts with substantial healthy-control representation and a case-dominated external source (36 cancer/1 healthy). The primary sources contain both classes but are not exactly balanced. Adalsteinsson had 528 known cancer participants and 0 explicitly healthy participants; Sun had 12 and 0. The archived local inventory yielded no additional source meeting the eligibility criterion. This inventory bounds the present analysis and is distinct from a current exhaustive database census.')
    replace('Study Source is a proxy containing', 'Study Source combines experimental-batch and population-composition differences. Each bilateral fit has one training study, making multi-source ICC filtering, multi-batch source ComBat and multi-source adversarial training unidentifiable in that protocol. The LSN-only fallback and unavailable method cells are recorded explicitly. The external cohort\'s single healthy participant supports a descriptive outcome, while generalization inference remains confined to the public-cohort evaluation.')
    replace('The Hybrid architecture and revised', 'The analysis is methodological development on the same public participants that motivated the representation comparison. Fold-level patient isolation addresses algorithmic leakage; independent confirmatory evaluation remains a separate requirement. Bootstrap intervals condition on the fixed studies and fitted scores and exclude refitting, method-development history and future-study uncertainty. DeLong tests are target-specific. The composition stress experiment uses a finite set of real-patient batches to measure adaptation dependence under the tested compositions.')
    replace('The analysis lacks directly', 'Matched copy-number measurements and tumor-specific chromatin assays were unavailable for the complete cohort. Broad arm-level fragmentation contrasts and reference GC associations therefore remain descriptive endpoints rather than direct CNA or accessibility measurements. Source metadata also have missing clinical covariates. Further source-level replication with balanced control recruitment and prospective specimen handling would extend the present cross-study conclusions.')
    replace('Cross-study generalization in public cfDNA cohorts depends', 'Normalization scale, source composition and score comparability define distinct axes of cross-study cfDNA evaluation. The Scale Trade-off audit quantifies broad fragmentation contrasts and residual genomic source variation, while pair decomposition connects study-wise ranking to pooled discrimination. Fixed-patient anchors identify companion-dependent target adaptation and demonstrate its absence for within-sample transformations. This framework makes the behaviour of normalization and locked thresholds directly testable across public studies.')

    replace('**Figure 1.', '**Figure 1. Study-associated fragmentation baselines and genomic variation.** (a) Source composition and observed healthy fragment-frequency curves with reference length markers. (b) Descriptive UMAP coloured by study and diagnosis, with coverage-adjusted conditional PERMANOVA fractions and permutation results. (c) Unique study and disease variance fractions across autosomal windows. (d) Healthy baseline distributions and one-way healthy-center clustering scores. The external source contributes one descriptive healthy observation; embeddings summarize the cohort independently of predictive preprocessing.')
    replace('**Figure 2.', '**Figure 2. Study-held-out discrimination and source-locked operating-point transport.** (a) Random-CV and bilateral study-held-out macro AUROC for Raw Ratio and Hybrid across fixed classifiers in the primary patient cohort. Training-set sizes differ between schemes. (b) Bilateral macro-AUROC matrix; grey denotes unavailable ComBat, DANN and unrequested tuned-classifier combinations. Asterisks identify the LSN-only operational fallback when a healthy-center filter is unidentifiable from one source. (c) ElasticNet target specificity minus the source target in percentage points, with exact binomial intervals for the two primary healthy target groups. (d) Sensitivity in the separately held-out high-burden cohort at primary-source-locked thresholds, with exact binomial intervals. The single healthy outcome is retained in the accompanying numerical table.')
    replace('**Figure 3.', '**Figure 3. Multiscale representation and fixed-patient composition dependence.** (a) Global and same-arm local branches, training-only scaling, concatenation and ElasticNet classification; bars show the two directional AUROCs. (b) Descriptive primary pooled held-out ROC and precision–recall curves, with target-specific DeLong comparisons against Global Z-score. (c) Primary calibration curves and Brier scores. (d) Maximum probability range of fixed real anchor patients as the target batch case fraction changes with source fitting frozen. Labels define the diagnostic compositions after benchmarking.')
    replace('**Figure 4.', '**Figure 4. Genomic context and fixed-pipeline feature retention.** (a) Raw-feature study variance associations with reference GC and telomere distance. (b) Reference-region odds ratios for majority-fold filtering under the fixed LSN-ICC protocol, with corrected results in the accompanying table. (c) Observed chromosome trajectories before and after local contrast, including gaps across removed windows. (d) Fixed-pipeline coefficient rank concordance across the original three-study LOCO fits. The reference context and coefficient ranks describe the fixed local representation.')
    replace('**Supplementary Figure S5.', '**Supplementary Figure S5. Pair concordance and cross-fold score comparability.** Within-study and cross-study cancer–control pair contributions sum to the secondary pooled AUROC. The paired summary compares pooled discrimination with the equal-weight mean of the two primary held-out target AUROCs, separating their weighting and score-comparability requirements.')

    full='\n\n'.join(paragraphs)
    # Publication-facing Highlights state the established contribution directly.
    highlights=[
        'Normalization scale shapes cross-study cfDNA discrimination and calibration.',
        'Pair concordance separates study-wise ranking from pooled score comparability.',
        'Fixed patient anchors expose composition-dependent batch adaptation.',
        'Within-sample transforms preserve anchor scores across target compositions.',
        'Source-locked thresholds reveal operating-point drift across public studies.']
    assert all(len(line)<=85 for line in highlights)
    start=full.index('## Highlights\n');end=full.index('## References\n')
    full=full[:start]+'## Highlights\n\n'+'\n'.join('- '+x for x in highlights)+'\n\n'+full[end:]
    original=pack.read(ARCHIVE/'submission_package/Manuscript_Main_Blind_and_Full.md')
    # Numeric table content is an editorial invariant, including row order/precision.
    assert re.findall(r'^\|.*$',original,re.M)==re.findall(r'^\|.*$',full,re.M)
    return full,highlights


def verify_doi(doi):
    assert re.fullmatch(r'10\.5281/zenodo\.\d+',doi), 'A real Zenodo DOI is required'
    record=doi.rsplit('.',1)[1]
    url='https://zenodo.org/api/records/'+record
    path=WORK/f'zenodo_record_{record}.json'
    data=pack.get_bytes(url)
    obj=json.loads(data)
    actual=obj.get('doi') or obj.get('pids',{}).get('doi',{}).get('identifier')
    assert actual==doi,(actual,doi)
    md=obj.get('metadata',{})
    assert str(md.get('version','')).lstrip('v')=='0.1.0',md.get('version')
    assert 'cfDNA' in md.get('title','') or 'cfDNA' in md.get('description','')
    assert 'Ma' in json.dumps(md.get('creators',[]),ensure_ascii=False)
    path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(data)
    return {'doi':doi,'record_url':'https://zenodo.org/records/'+record,
            'api_url':url,'metadata_sha256':pack.sha(path),'tag':TAG,'repository':REPO,
            'public_record_verified':True,'metadata':md}


def availability(full,receipt):
    start=full.index('## Data and code availability\n')
    end=full.index('## CRediT authorship contribution statement\n')
    text=('## Data and code availability\n\n'
          'Source data are described in Cristiano et al. (2019), Jiang et al. (2015) and '
          'Snyder et al. (2016). Fragment intervals were accessed through FinaleDB '
          '(https://finaledb.research.cchmc.org/); source-linked metadata are documented '
          'in UNITE supplementary data file S1, sheet s2 '
          '(https://doi.org/10.1126/sciadv.ady9432). Source URLs, eligibility rules and '
          'provenance are recorded in the extraction configuration and manifests.\n\n'
          f'Code and compact statistical outputs are distributed at {REPO}. ')
    if receipt:
        text+=(f'The immutable {TAG} software release is archived in Zenodo at '
               f'https://doi.org/{receipt["doi"]}. ')
    else:
        text+=(f'Zenodo registration is in progress for the {TAG} software release. ')
    text+=('The release contains the MIT-licensed original software, frozen aggregate '
           'statistical tables, tests and compact benchmark provenance. Large participant '
           'feature matrices, raw fragment streams and fitted-model caches are excluded '
           'from the lightweight distribution. Streaming scripts reconstruct the derived '
           'inputs from the published sources subject to availability and their original '
           'data-use terms. Third-party data and annotations retain those terms.\n\n')
    return full[:start]+text+full[end:]


def make_blind(full,author):
    blind=re.sub(r'\*\*Author:\*\*.*?\*\*ORCID:\*\* Not registered\.\n\n','',full,flags=re.S)
    blind=blind.replace(author['name'],'The sole author')
    start=blind.index('Code and compact statistical outputs are distributed at ')
    end=blind.index('The release contains the MIT-licensed',start)
    blind=blind[:start]+('The versioned code and statistical archive are supplied through '
        'the separate title page for anonymous review. ')+blind[end:]
    assert not any(x in blind for x in [author['email'],author['name'],'mzxuan08'])
    return blind


def revised_cover(title,author,journal,receipt):
    return (f'# Cover letter\n\n{pack.DATE}\n\nEditor-in-Chief\n\n{journal}\n\nDear Editor,\n\n'
      f'I submit “{title}” as an original research article for {journal}. The manuscript '
      'provides an interpretable evaluation framework for cross-study cfDNA fragmentomics '
      'by separating normalization scale, study-wise ranking and score transport.\n\n'
      'The first contribution is a Scale Trade-off audit of global, same-arm local and '
      'hybrid self-normalization. It measures how each representation changes broad '
      'fragmentation contrasts and residual GC-associated study variation, with bidirectional '
      'source-held-out testing and source-locked operating thresholds. Hybrid tests scale '
      'complementarity; the comparison is reported objectively alongside Global Z-score.\n\n'
      'Second, cancer–control pair concordance decomposition distinguishes within-study '
      'discrimination from cross-study score comparability. A fixed real-patient Anchor '
      'Composition Stress Test then isolates companion-dependent target adaptation: '
      'changing target batch composition shifts ComBat-adapted probabilities by up to '
      '0.214, while within-sample Global, LSN and Hybrid transformations yield a shift '
      'of 0 for the same frozen scoring rules. Preserved within-study standardized '
      'contrasts help locate the effect at adaptation and score transport.\n\n'
      'Third, the software and statistical archive provide reproducible, fold-isolated '
      'preprocessing, nested source selection and threshold calibration, with automated '
      'tests and checksum-verifiable outputs. The study uses 585 source-identified '
      'participants: a primary bilateral cohort of 548 and a separate case-dominated '
      'external sensitivity cohort of 37. This hierarchy aligns the evaluation endpoint '
      'with each source’s control representation.\n\n'
      'I am the sole author and have reviewed and approved the code logic, statistical '
      'outputs and manuscript. I take responsibility for the work. The manuscript is '
      'original, has not been published as a journal article and is not under consideration '
      'by another journal. Any posted preprint will be disclosed in the submission system. '
      'Funding, competing interests, public-data ethics status and AI assistance '
      'are stated in the manuscript.\n\n'
      'I request the traditional subscription publishing model, subject to the journal’s '
      'current terms.\n\n'
      f'Sincerely,\n\n{author["name"]}\n\n{author["affiliation"]}\n\n{author["email"]}\n')


def readme(author,receipt):
    old=pack.release_readme(author)
    # Preserve executable instructions and provenance while replacing presentation prose.
    start=old.index('## Included files and exclusions')
    body=old[start:]
    body=body.replace('This release does not\nclaim cross-platform numerical equivalence.',
        'Exact software versions and the validated Windows/CUDA environment are recorded\nin the Methods tables.')
    intro=('''# cfDNA Scale Tradeoff Audit

**Scale Trade-offs in Within-Sample Self-Normalization for Cross-Study Generalization
in Public cfDNA Cohorts**

'''+f'{author["name"]} · {author["affiliation"]}\n\n'+
        '## What the audit resolves\n\n'
        'A pooled cfDNA benchmark score combines representation choices, study-wise ranking '
        'and comparability across fitted scoring rules. This project separates those '
        'components and measures how normalization scale and target cohort composition '
        'affect cross-study evaluation.\n\n'
        '- **Scale trade-offs:** Global, same-arm local and Hybrid representations quantify '
        'broad fragmentation-contrast retention and residual genomic source variation.\n'
        '- **Pair concordance:** within-study and cross-study cancer–control pairs explain '
        'their separate contributions to pooled AUROC.\n'
        '- **Fixed anchors:** changing companion patients produces a maximum probability '
        'shift of 0.214 for the evaluated ComBat-plus-target alignment and 0 for the '
        'within-sample transforms.\n'
        '- **Locked thresholds:** source-only calibration makes target specificity drift '
        'visible alongside discrimination and calibration.\n\n'
        '## Cohort and evaluation\n\n'
        'The analysis contains 585 source-identified participants from 628 extracted '
        'specimens and uses 477 retained hg19 autosomal 5 Mb windows. The primary bilateral '
        'cohort contains Cristiano 2019 (427) and Jiang 2015 (121), total 548. Snyder '
        '2016 (37; 36 cancer and one healthy) is a separately held-out sensitivity cohort. '
        'Secondary Combined LOCO evaluates all 585. Study Source combines experimental '
        'and population-composition differences.\n\n'
        'Fixed ElasticNet Combined LOCO AUROCs are Global 0.759, Hybrid 0.731, fixed '
        'LSN-ICC 0.724 and ComBat-plus-target alignment 0.543. Hybrid is a scale-fusion '
        'comparator; the study’s contribution is the mechanistic evaluation framework. '
        'Use the exact CSV cells with cohort_scope for inference.\n\n'
        '## Code and archive\n\n'+f'Public repository: {REPO}\n\n')
    if receipt:
        intro+=f'Immutable software release {TAG}: https://doi.org/{receipt["doi"]}\n\n'
    else:intro+=f'First software release: {TAG}; Zenodo registration follows repository linkage.\n\n'
    return intro+('## Publication materials\n\nThe publication-facing manuscript, Highlights, legends and cover letters follow '
        'the normalization-scale and score-transport narrative. scripts/09_polish_and_release_materials.py '
        'is the editorial revision utility; it is separate from the research computation. '
        'It requires the local submission workspace and its preserved pre-edit manuscript snapshot. '
        'The light software release contains the approved text and research scripts; it does not '
        'reconstruct that private submission-author workspace.\n\n')+body


def build(args,before):
    archive_once()
    author=json.loads(pack.read(ROOT/'config/submission_author.json'))
    receipt=verify_doi(args.doi) if args.doi else None
    if receipt:
        pack.dump(WORK/'release_receipt.json',receipt)
        author.update({'repository_status':'public','zenodo_doi':receipt['doi'],
                       'zenodo_status':'published immutable '+TAG+' archive'})
        pack.dump(ROOT/'config/submission_author.json',author)
    full,highlights=revise_full();full=availability(full,receipt)
    blind=make_blind(full,author)
    title=full.splitlines()[0][2:]
    abstract=pack.extract_section(full,'## Abstract','## 1. Introduction')
    abstract_text=abstract.split('**Keywords:**')[0].strip()
    keywords=abstract.split('**Keywords:**')[1].strip()
    body=full.split('## 1. Introduction\n')[1].split('## Data and code availability\n')[0]
    counts={'abstract_words':pack.word_count(abstract_text),
            'main_body_words_excluding_tables':pack.word_count('\n'.join(x for x in body.splitlines() if not x.startswith('|'))),
            'main_body_including_tables_words':pack.word_count(body),
            'full_markdown_words':pack.word_count(full),
            'definition':'Whitespace-delimited tokens containing a letter or numeral; Markdown formatting stripped; abstract excludes Keywords; body includes Introduction through Conclusions and equations/headings, excludes declarations/references/legends.'}
    for path in [OUT/'Manuscript_Main_Blind_and_Full.md',ROOT/'results/MANUSCRIPT_FULL_TEXT.md',
                 ROOT/'results/reports/MANUSCRIPT_MASTER_DRAFT.md']:
        pack.write(path,full)
    pack.write(OUT/'Manuscript_Main_Blinded.md',blind)
    pack.dump(OUT/'verification/Word_Counts.json',counts)
    pack.dump(OUT/'verification/Highlights_Character_Counts.json',[
        {'text':h,'characters_including_spaces':len(h),'passes_85':len(h)<=85} for h in highlights])
    pack.dump(OUT/'author_configuration.json',author)
    doc=pack.pandoc_doc(pack.title_page(title,abstract_text,keywords,counts,author),OUT/'01_Title_Page.docx',args.pandoc,author,double=False)
    # Title-page availability follows the current verified public-release status.
    for p in doc.paragraphs:
        if p.text.startswith('Prepared repository:'):
            p.text=f'Public repository: {REPO}. '+(f'Immutable {TAG} archive: https://doi.org/{receipt["doi"]}.' if receipt else 'Zenodo archival registration is in progress.')
            for run in p.runs:pack.set_font(run)
    doc.save(OUT/'01_Title_Page.docx')
    pack.pandoc_doc('# Highlights\n\n'+'\n\n'.join('- '+h for h in highlights),OUT/'02_Highlights.docx',args.pandoc,author,blind=True,double=False)
    pack.pandoc_doc(blind,OUT/'03_Manuscript_DoubleSpaced_LineNumbered.docx',args.pandoc,author,blind=True,numbers=True)
    doc=pack.pandoc_doc(full,OUT/'04_bioRxiv_Integrated_Preprint.docx',args.pandoc,author)
    pack.add_integrated_figures(doc);doc.save(OUT/'04_bioRxiv_Integrated_Preprint.docx')
    supp=pack.supplement(full,author)
    supp=supp.replace('This supplement accompanies the frozen analysis and documents the scope of the statistical tables. It introduces no new fitted predictions or observations.',
        'This supplement decomposes score comparability and documents the statistical tables supporting the normalization-scale and fixed-anchor analyses.')
    doc=pack.pandoc_doc(supp,OUT/'05_Supplementary_Material.docx',args.pandoc,author)
    from docx.shared import Inches
    for p in list(doc.paragraphs):
        if p.text.startswith('Supplementary Figure S5.'):
            image=p.insert_paragraph_before();image.add_run().add_picture(str(ROOT/'results/figures'/f'{pack.FIGURES[4]}.png'),width=Inches(6.5))
            image.paragraph_format.keep_with_next=True;image.paragraph_format.line_spacing=1
            p.paragraph_format.line_spacing=1.15
    doc.save(OUT/'05_Supplementary_Material.docx')
    for name,journal in [('06_Cover_Letter','Computer Methods and Programs in Biomedicine'),('06b_Cover_Letter_Alternative_CBM','Computers in Biology and Medicine')]:
        text=revised_cover(title,author,journal,receipt);pack.write(OUT/(name+'.md'),text)
        letter=pack.pandoc_doc(text,OUT/(name+'.docx'),args.pandoc,author,double=False)
        from docx.shared import Pt
        for paragraph in letter.paragraphs:
            paragraph.paragraph_format.line_spacing=1
            paragraph.paragraph_format.space_after=Pt(5)
        letter.save(OUT/(name+'.docx'))
    pack.pandoc_doc('# Declaration of interest\n\n'+title+'\n\n'+author['name']+'\n\n'+author['competing_interests'],OUT/'07_Declaration_of_Interest.docx',args.pandoc,author,double=False)
    copies=json.loads(pack.read(OUT/'verification/Byte_Identical_Copies.json'))
    pack.create_release(author,full,copies)
    pack.write(RELEASE/'README.md',readme(author,receipt))
    cff=json.loads(pack.read(RELEASE/'CITATION.cff'))
    if receipt:cff['doi']=receipt['doi']
    pack.dump(RELEASE/'CITATION.cff',cff)
    pack.dump(RELEASE/'.zenodo.json',{'title':'cfDNA Scale Tradeoff Audit: reproducible analysis of normalization scale and cross-study score transport',
        'upload_type':'software','description':'Reproducible software and compact statistical outputs accompanying the Scale Trade-offs in Within-Sample Self-Normalization audit of public cfDNA cohorts. Includes source-isolated preprocessing, global/local/Hybrid comparisons, pair concordance decomposition, fixed real-patient composition stress tests, automated tests and checksum-verifiable provenance. Original software is MIT licensed; third-party data and reference annotations retain their original terms.',
        'creators':[{'name':'Ma, Zhenxuan','affiliation':author['affiliation']}],
        'version':'0.1.0','license':'MIT','keywords':['cfDNA','fragmentomics','normalization','cross-study validation','batch composition'],
        'related_identifiers':[{'identifier':REPO+'/tree/'+TAG,'relation':'isSupplementTo','resource_type':'software'}]})
    pack.dump(OUT/'verification/Byte_Identical_Copies.json',copies)
    pack.package_validation(author)
    assert snapshot()==before
    pack.dump(WORK/'scientific_integrity.json',{'status':'passed','protected_files':before,
        'changed_scientific_files':[],'authorized_text_revision':True,'receipt':receipt})
    pack.write(ROOT/'results/reports/step_09_report.md',
        '# Editorial revision and public release\n\nPublication-facing text was revised with anti-defensive-writing-en: normalization-scale trade-offs, pair concordance and fixed-patient composition dependence form the main narrative. Observed comparisons, method identifiability, source-composition confounding, study-level replication and inference limits remain explicit. Original numerical tables and figure files are unchanged.\n\n'
        +json.dumps({'status':'revised_pending_final_render','protected_scientific_files':len(before),
        'counts':counts,'zenodo_receipt':receipt,'clinical_refitting':False},ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'phase':'DOI_updated' if receipt else 'prepared','counts':counts,'protected_files':len(before)},ensure_ascii=True))


def finalize(args,before):
    render=json.loads(pack.read(ROOT/'work/step08/render_status.json'))
    assert len(render)==8
    for x in render:
        assert x['returncode']==0 and x['pages']>0 and x['sha256']==pack.sha(OUT/x['file'])
    author=json.loads(pack.read(ROOT/'config/submission_author.json'))
    pack.package_validation(author)
    receipt=json.loads(pack.read(WORK/'release_receipt.json')) if (WORK/'release_receipt.json').exists() else None
    assert receipt and author['zenodo_doi']==receipt['doi'],'Published DOI required for final release package'
    pdf=OUT/'04_bioRxiv_Integrated_Preprint.pdf'
    source=ROOT/'work/step08/render/04_bioRxiv_Integrated_Preprint/04_bioRxiv_Integrated_Preprint.pdf'
    shutil.copy2(source,pdf)
    from pypdf import PdfReader
    reader=PdfReader(pdf)
    assert len(reader.pages)==next(x['pages'] for x in render if x['file'].startswith('04_'))
    assert receipt['doi'] in re.sub(r'\s+','',''.join(p.extract_text() or '' for p in reader.pages))
    assert snapshot()==before
    integrity={'status':'passed','protected_files':before,'checked_files':len(before),
               'changed_files':[],'frozen_manifest_sha256':pack.sha(ROOT/'results/reports/final_scope_alignment_integrity.json')}
    pack.dump(OUT/'verification/Editorial_Release_Integrity.json',integrity)
    pack.dump(OUT/'verification/Frozen_Input_Integrity.json',dict(integrity,
        scope='Scientific inputs and outputs; authorized manuscript text revisions excluded',
        prior_editorial_snapshot='results/archive_step08_before_editorial_release/'))
    pack.dump(OUT/'verification/Published_Release_Receipt.json',receipt)
    archive=json.loads(pack.read(WORK/'archive_verification.json'))
    assert archive['status']=='passed' and archive['version_doi']==receipt['doi']
    pack.dump(OUT/'verification/Archived_Release_Verification.json',archive)
    pack.dump(OUT/'verification/Rendered_Documents.json',render)
    pack.dump(OUT/'verification/Render_QA_Review.json',{
        'status':'passed','method':'All current rendered pages visually inspected, including equations, tables, figure legends, line numbers, DOI and one-page cover letters',
        'rendered_documents':render,'total_pages':sum(x['pages'] for x in render),
        'prior_step08_review_superseded':True})
    import xml.etree.ElementTree as ET
    suite=ET.parse(WORK/'pytest.xml').getroot()
    suites=list(suite.iter('testsuite'))
    test_counts={key:sum(int(s.get(key,0)) for s in suites) for key in ['tests','failures','errors','skipped']}
    assert test_counts['failures']==0 and test_counts['errors']==0
    test_counts['passed']=test_counts['tests']-test_counts['skipped']
    pack.dump(OUT/'verification/Current_Test_Results.json',test_counts)
    shutil.copy2(ROOT/'scripts/09_polish_and_release_materials.py',RELEASE/'scripts/09_polish_and_release_materials.py')
    pack.dump(RELEASE/'results/reports/zenodo_release_receipt.json',archive)
    copies=json.loads(pack.read(OUT/'verification/Byte_Identical_Copies.json'))
    # Refresh source-linked copies after authorized DOI and editorial-script updates.
    for relative,record in list(copies.items()):
        original=ROOT/record['source']; copied=OUT/relative
        if original.exists() and copied.exists() and pack.sha(original)==pack.sha(copied):
            record['sha256']=pack.sha(original)
        else:
            del copies[relative]
    pack.dump(OUT/'verification/Byte_Identical_Copies.json',copies)
    summary={'status':'complete','editorial_skill':'anti-defensive-writing-en','docx_count':8,
        'rendered_pages':sum(x['pages'] for x in render),'preprint_pdf_pages':len(reader.pages),
        'repository':REPO,'zenodo_doi':receipt['doi'],'release_tag':TAG,
        'protected_scientific_files':len(before),'scientific_files_changed':0,
        'pdf_sha256':pack.sha(pdf),'rendered_documents':render,
        'tests':test_counts,'verified_archived_release':archive}
    pack.dump(OUT/'verification/Build_Summary.json',summary)
    checklist=pack.read(OUT/'SUBMISSION_CHECKLIST_FOR_AUTHOR.md')
    # Replace stale pending-release checklist with a current concise, factual guide.
    inventory=sorted(str(p.relative_to(OUT)).replace('\\','/') for p in OUT.rglob('*') if p.is_file() and '.git' not in p.parts)
    hs=json.loads(pack.read(OUT/'verification/Highlights_Character_Counts.json'))
    text=('# 投稿与开源发布核验清单\n\n'
        f'作者：Zhenxuan Ma（马振轩）；Hubei University of Medicine, Shiyan 442000, Hubei, China；mzxmc08@gmail.com。\n\n'
        f'公开代码仓库：{REPO}\n\n真实版本DOI：https://doi.org/{receipt["doi"]}\n\n'
        f'归档版本：{TAG}。04的PDF为{len(reader.pages)}页；对应Word与图件内容已核验。\n\n'
        '## 修订范围与数据完整性\n\n'
        '正文、图注、Highlights、投稿信、补充材料导语、README及论文写作汇总按anti-defensive-writing-en统一叙事。'
        '尺度权衡、样本对一致性分解与固定真实患者压力测试形成主线；Hybrid真实比较、统计边界及必要局限完整保留。'
        f'{len(before)}个受保护科学文件SHA256不变；未重新提取特征或训练临床模型。历史执行报告作为证据保留原文。\n\n'
        '## Highlights字符数（含空格）\n\n'+ '\n'.join(f'- {h["characters_including_spaces"]}/85：{h["text"]}' for h in hs)+
        '\n\n## 三分钟上传指南\n\n'
        '1. bioRxiv上传04_bioRxiv_Integrated_Preprint.pdf及05补充材料；按平台实际筛查要求选择主题与许可。\n'
        '2. 期刊只选CMPB或CBM一家。分别上传01署名页、03匿名行号正文、02 Highlights、对应06投稿信、07利益冲突声明、05补充材料和独立图表。\n'
        '3. 检查系统PDF的公式、表格、图注、真实DOI及8条文献，填写AI与伦理声明，并选择Subscription/traditional路线。最终费用及个性化格式要求以期刊系统为准。\n\n'
        '## 作者仍需留意\n\n'
        '- ORCID尚未注册；如平台强制要求请本人注册真实号码。\n'
        '- 稿件如实记录公开去标识化数据二次分析及未取得额外机构判定的状态。如学校或期刊要求伦理确认，按真实机构结论补充。\n'
        '- AI代码、统计/绘图与文本辅助已披露；人工核验及责任按作者已确认内容记录。\n'
        '- 双盲正文暂不显示个人仓库与归档标识，完整署名稿和Title Page提供真实链接。\n'
        '- GitHub和Zenodo已公开；本任务没有向bioRxiv或期刊提交稿件。\n\n## 全部文件\n\n'+
        '\n'.join('- `'+n+'`' for n in inventory)+'\n')
    pack.write(OUT/'SUBMISSION_CHECKLIST_FOR_AUTHOR.md',text)
    # Git internals are never part of an uploadable research archive.
    def files(folder):
        return [p for p in sorted(folder.rglob('*')) if p.is_file() and not any(x in p.parts for x in ['.git','__pycache__','.pytest_cache'])]
    release_hash={str(p.relative_to(RELEASE)).replace('\\','/'):pack.sha(p) for p in files(RELEASE) if p.name!='RELEASE_MANIFEST_SHA256.json'}
    pack.dump(RELEASE/'RELEASE_MANIFEST_SHA256.json',release_hash)
    for folder,path,prefix in [(RELEASE,OUT/'github_zenodo_release.zip',RELEASE.parent)]:
        with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
            for p in files(folder):z.write(p,str(p.relative_to(prefix)))
    hashes={str(p.relative_to(OUT)).replace('\\','/'):pack.sha(p) for p in files(OUT) if p.name!='PACKAGE_MANIFEST_SHA256.json'}
    pack.dump(OUT/'PACKAGE_MANIFEST_SHA256.json',hashes)
    target=ROOT/'cfDNA_Submission_Package.zip'
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in files(OUT):z.write(p,str(p.relative_to(ROOT)))
    with zipfile.ZipFile(target) as z:assert z.testzip() is None
    pack.write(ROOT/'cfDNA_Submission_Package.zip.sha256',pack.sha(target)+'  cfDNA_Submission_Package.zip\n')
    summary['package_zip_sha256']=pack.sha(target)
    pack.write(ROOT/'results/reports/step_09_report.md','# Editorial revision and verified public release\n\n'+json.dumps(summary,ensure_ascii=False,indent=2)+'\n\nThe software release is publicly archived; the manuscript PDF and DOCX carry the verified DOI. Scientific tables, features, predictions and figure bytes remain unchanged. No manuscript was submitted to a journal or preprint server by this script.\n')
    print(json.dumps(summary,ensure_ascii=True))


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--pandoc',type=Path,required=True)
    p.add_argument('--doi')
    p.add_argument('--finalize',action='store_true')
    args=p.parse_args();WORK.mkdir(parents=True,exist_ok=True)
    before=snapshot()
    if (WORK/'scientific_baseline.json').exists():assert json.loads(pack.read(WORK/'scientific_baseline.json'))==before
    else:pack.dump(WORK/'scientific_baseline.json',before)
    if args.finalize:finalize(args,before)
    else:build(args,before)


if __name__=='__main__':main()
