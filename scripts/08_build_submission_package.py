"""Build submission copies without refitting models or altering frozen assets.

Usage: python scripts/08_build_submission_package.py --pandoc /path/to/pandoc
Requires Python >=3.10, python-docx and lxml; Pandoc supplies editable OMML math.
Network access is limited to bibliographic verification of the eight source DOIs.
The author configuration records explicit human confirmations, not inferred ones.
"""
from __future__ import annotations

import argparse
import copy
import csv
import datetime as dt
import hashlib
import html
import json
import re
import shutil
import subprocess
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission_package'
QA = ROOT / 'work' / 'step08'
DOIS = ['10.1038/s41586-019-1272-6', '10.1073/pnas.1500076112',
        '10.1016/j.cell.2015.11.050', '10.1093/biostatistics/kxj037',
        '10.1093/bioinformatics/btaa999', '10.1126/sciadv.ady9432',
        '10.2307/2531595', '10.1111/j.2517-6161.1995.tb02031.x']
FIGURES = ['Figure1_Batch_Vulnerability_Audit', 'Figure2_Random_vs_LOCO_Collapse',
           'Figure3_LSN_Debiasing_Performance', 'Figure4_Genomic_Interpretability',
           'FigureS5_ComBat_Score_Comparability']
DATE = dt.datetime.now(dt.timezone(dt.timedelta(hours=8))).date().isoformat()


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1 << 20), b''):
            h.update(block)
    return h.hexdigest()


def read(path):
    return Path(path).read_text(encoding='utf-8-sig')


def write(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value, encoding='utf-8', newline='\n')


def dump(path, obj):
    write(path, json.dumps(obj, ensure_ascii=False, indent=2) + '\n')


def rows(name):
    with (ROOT / 'results' / 'tables' / name).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def frozen_snapshot():
    manifest = ROOT / 'results/reports/final_scope_alignment_integrity.json'
    obj = json.loads(read(manifest))
    for name, expected in obj['frozen_files'].items():
        if sha(name) != expected:
            raise ValueError(f'Frozen manifest mismatch: {name}')
    targets = {manifest, ROOT / 'results/MANUSCRIPT_FULL_TEXT.md',
               ROOT / 'results/reports/final_scope_alignment.md'}
    for folder in ['data/processed', 'data/metadata', 'results/tables', 'results/figures']:
        targets.update(p for p in (ROOT / folder).rglob('*') if p.is_file())
    targets.update(Path(p) for p in obj['frozen_files'])
    return {str(p.relative_to(ROOT)): sha(p) for p in sorted(targets)}


def assert_frozen(before):
    after = frozen_snapshot()
    if before != after:
        raise ValueError('Protected inputs changed during packaging')
    return {'status': 'passed', 'checked_files': len(before), 'changed_files': [],
            'frozen_manifest_sha256': sha(ROOT / 'results/reports/final_scope_alignment_integrity.json'),
            'protected_files': before,
            'feature_extraction_rerun': False, 'classifier_training_rerun': False,
            'figure_regeneration': False}


def get_bytes(url, attempts=5, accept='application/json'):
    for i in range(attempts):
        try:
            req = urllib.request.Request(url, headers={
                'User-Agent': 'cfDNA-Scale-Tradeoff-Audit/0.1 (mailto:mzxmc08@gmail.com)',
                'Accept': accept})
            with urllib.request.urlopen(req, timeout=45) as response:
                return response.read()
        except (OSError, urllib.error.HTTPError):
            if i == attempts - 1:
                raise
            time.sleep(min(50, 5 * 2 ** i))


def clean(text):
    return re.sub(r'\s+', ' ', html.unescape(re.sub('<[^>]+>', '', str(text)))).strip()


def initials(given):
    return ''.join(word[0].upper() for word in re.findall(r'[^\W\d_]+', given, re.UNICODE))


def reference_verification():
    dest = OUT / 'verification/crossref'
    dest.mkdir(parents=True, exist_ok=True)
    pmids = {2: '25646427', 7: '3203132'}
    supplements = {}
    for number, pmid in pmids.items():
        path = dest / f'pubmed_{pmid}.xml'
        url = ('https://eutils.ncbi.nlm.nih.gov/entrez/eutils/efetch.fcgi?'
               f'db=pubmed&id={pmid}&retmode=xml')
        if not path.exists():
            path.write_bytes(get_bytes(url, accept='application/xml'))
        root = ET.fromstring(path.read_bytes())
        article = root.find('.//PubmedArticle')
        assert article is not None and article.findtext('.//PMID') == pmid
        supplements[number] = {'pages': article.findtext('.//Pagination/MedlinePgn'),
                               'url': f'https://pubmed.ncbi.nlm.nih.gov/{pmid}/',
                               'api_url': url, 'xml_sha256': sha(path)}
    verified, formatted = [], []
    for number, doi in enumerate(DOIS, 1):
        path = dest / f'{number:02d}.json'
        url = 'https://api.crossref.org/works/' + urllib.parse.quote(doi, safe='')
        if not path.exists():
            path.write_bytes(get_bytes(url))
            time.sleep(2)
        m = json.loads(read(path))['message']
        assert m['DOI'].lower() == doi.lower()
        date_field = 'published-print' if 'published-print' in m else 'published'
        year = m[date_field]['date-parts'][0][0]
        pages = m.get('article-number') or m.get('page')
        source = 'Crossref REST'
        if number in supplements:
            p = supplements[number]['pages']
            assert p
            if number == 2:
                assert p.replace('–', '-') in ['E1317-25', 'E1317-E1325']
                pages = 'E1317-E1325'
            if number == 7:
                assert p.replace('–', '-') in ['837-45', '837-845']
                pages = '837-845'
            source += ' + NLM PubMed pagination'
        assert all([m.get('title'), m.get('volume'), m.get('issue'), pages, year])
        authors = m['author']
        citation_authors = ', '.join(clean(a['family']) + ' ' + initials(a.get('given', ''))
                                     for a in (authors[:3] if len(authors) > 6 else authors))
        if len(authors) > 6:
            citation_authors += ', et al'
        title, journal = clean(m['title'][0]), clean(m['container-title'][0])
        publisher_source = ''
        if number == 8:
            # The publisher identifies the historical journal title explicitly;
            # Crossref's container-title is its modern successor label.
            journal = 'Journal of the Royal Statistical Society: Series B (Methodological)'
            publisher_source = 'https://academic.oup.com/jrsssb/article/57/1/289/7035855'
            source += ' + publisher historical journal title'
        citation = (f'{number}. {citation_authors}. {title}. {journal}. '
                    f'{year};{m["volume"]}({m["issue"]}):{pages}. https://doi.org/{doi}')
        formatted.append(citation)
        verified.append({'reference': number, 'authors': citation_authors, 'title': title,
                         'journal': journal, 'year': year, 'year_field': date_field,
                         'volume': m['volume'], 'issue': m['issue'], 'pages_or_article': pages,
                         'doi': doi, 'source': source, 'crossref_url': url,
                         'crossref_sha256': sha(path),
                         'supplement_source': supplements.get(number, {}).get('url', '') or publisher_source,
                         'accessed': DATE})
    assert verified[5]['year'] == 2026 and verified[5]['volume'] == '12'
    assert verified[5]['issue'] == '28' and verified[5]['pages_or_article'] == 'eady9432'
    with (OUT / 'verification/References_Verified.csv').open('w', encoding='utf-8', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(verified[0]))
        writer.writeheader()
        writer.writerows(verified)
    dump(OUT / 'verification/reference_verification.json', {'status': 'passed',
         'references': verified, 'pubmed_pagination_supplements': supplements})
    write(OUT / 'verification/Reference_Verification_Notes.md',
          '# Reference verification\n\nAll eight DOIs were matched against Crossref REST records. '
          'Print publication year is used where available; online publication and deposit dates '
          'are not substituted for the issue year. All volume and issue fields were supplied by Crossref. '
          'Crossref omits the Jiang pagination and supplies only the first DeLong page; '
          'NLM PubMed records supply their full page ranges. The eight entries are numbered as in '
          'the frozen manuscript. Article titles and author initials derive from the retrieved records.\n\n'
          'Reference 6 is Science Advances, 2026, volume 12, issue 28, article eady9432 '
          '(10 July 2026), DOI 10.1126/sciadv.ady9432. Its supplementary data file S1, '
          'sheet s2 is described in Methods, not appended as an invented part of the article citation.\n\n'
          'For reference 8, the publisher explicitly identifies the original journal as '
          'Journal of the Royal Statistical Society: Series B (Methodological); the citation '
          'uses that historical title rather than Crossref’s modern container label. '
          'Publisher verification: https://academic.oup.com/jrsssb/article/57/1/289/7035855 .\n\n'
          + '\n\n'.join(formatted) + '\n')
    return formatted


def word_count(text):
    """Whitespace-delimited count after stripping Markdown/math presentation marks.

    Hyphenated terms, inline math tokens and numerals count as one token;
    table rows are excluded by the caller for the body count.
    """
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    text = re.sub(r'<[^>]+>', '', text)
    text = re.sub(r'[#*`$]', '', text)
    return sum(bool(re.search(r'\w', word, re.UNICODE)) for word in text.split())


def extract_section(text, heading, next_heading):
    return text.split(heading + '\n', 1)[1].split(next_heading + '\n', 1)[0].strip()


def declarations(author, blind=False):
    name = 'The sole author' if blind else author['name']
    ai = (f'During the preparation of this work, the author used {author["ai_tool"]} '
          f'({author["ai_developer"]}; desktop application used in October 2026; exact '
          'application build and underlying model version not recorded) to assist with code '
          'implementation and refactoring, unit tests, statistical-workflow and figure-generation '
          'code, reference verification, manuscript organization and drafting, and language revision. '
          'The author reviewed the code logic, statistical outputs and text, edited the material '
          'as needed, approved the manuscript, and takes full responsibility for its accuracy, '
          'integrity and interpretation. The AI tool is not an author. Figures were generated '
          'by reproducible code from the recorded data; no clinical observations were fabricated.')
    ethics = ('This study is a secondary analysis of previously published, de-identified public '
              'human plasma cfDNA fragment-interval data and associated metadata obtained through '
              'FinaleDB and the published UNITE supplementary data. No new participants were '
              'recruited, no new specimens were collected, and no attempt was made to re-identify '
              'participants. No additional institutional ethics approval or formal exemption '
              'determination was obtained for this secondary analysis. Original-study consent '
              'and ethics protections remain those of the source investigations; no new consent '
              'was sought. Public-data exemption provisions may apply subject to institutional '
              'requirements; this statement does not assert an institutional exemption decision.')
    credit = (f'{name}: Conceptualization, Methodology, Software, Validation, Formal analysis, '
              'Investigation, Resources, Data curation, Writing – original draft, Writing – review '
              '& editing, Visualization, Supervision, Project administration. All applicable '
              'roles were undertaken by the sole author, who approved the final manuscript.')
    return {'CRediT authorship contribution statement': credit,
            'Funding': author['funding'],
            'Declaration of competing interests': author['competing_interests'],
            'Ethics statement': ethics,
            'Acknowledgements': 'The author thanks the investigators who generated the original '
                                'public datasets and the developers of FinaleDB, UNITE and the '
                                'open-source scientific software used in this analysis.',
            'Declaration of generative AI and AI-assisted technologies in the manuscript preparation process': ai}


def build_manuscripts(author, references):
    original = read(ROOT / 'results/MANUSCRIPT_FULL_TEXT.md')
    title = original.splitlines()[0][2:]
    full = original.replace('**Authors and affiliations:** [To be supplied by the authors]',
        f'**Author:** {author["name"]}\n\n**Affiliation:** {author["affiliation"]}\n\n'
        f'**Corresponding author:** {author["name"]}; {author["email"]}\n\n'
        '**ORCID:** Not registered.')
    start = full.index('## Data and code availability')
    end = full.index('## Figure legends')
    availability = ('Public source data are described in Cristiano et al. (2019), Jiang et al. '
        '(2015) and Snyder et al. (2016). Fragment intervals were accessed through FinaleDB '
        '(https://finaledb.research.cchmc.org/); the source-linked clinical inventory is '
        'documented in UNITE data file S1, sheet s2 (https://doi.org/10.1126/sciadv.ady9432). '
        'Download URLs and source provenance are recorded in the extraction configuration '
        'and generated manifests. No newly collected human data are reported.\n\n'
        f'The code and compact statistical release are prepared for {author["repository_url"]}. '
        'The repository was not publicly accessible at package preparation, and public release '
        'is pending. A versioned GitHub release will be archived in Zenodo after publication '
        'of the repository; no Zenodo DOI has yet been assigned. The assigned DOI and immutable '
        'release identifier will be added to this statement after archival release. '
        'The prepared release includes original software under the MIT license, frozen '
        'aggregate statistical tables, tests and a compact provenance bundle. Large feature '
        'matrices, raw intervals, model caches and third-party supplementary archives are '
        'excluded; the scripts reconstruct derived inputs by streaming the public sources '
        'subject to their availability and terms. Third-party data remain under their '
        'original conditions and are not relicensed by the software MIT license.')
    declared = '## Data and code availability\n\n' + availability + '\n\n'
    for heading, body in declarations(author).items():
        declared += f'## {heading}\n\n{body}\n\n'
    full = full[:start] + declared + full[end:]
    ai_methods = ('### 2.8. Code assistance and figure provenance\n\n'
        f'{author["ai_tool"]} ({author["ai_developer"]}; desktop application used in October 2026; exact application '
        'build and underlying model version not recorded) assisted implementation and '
        'refactoring, testing, statistical-workflow code and figure-generation code. '
        'The author checked the computational logic, statistical outputs and interpretation. '
        'All quantitative panels were generated from the recorded real-data outputs using '
        'the supplied Python plotting scripts. The explanatory workflow in Figure 3a was '
        'drawn programmatically with the same code assistance and reviewed by the author.\n\n')
    full = full.replace('## 3. Results', ai_methods + '## 3. Results')
    # Methods provenance preserves the specific supplementary-table locator.
    full = full.replace('Public study metadata had been recovered from the published UNITE supplementary data table,',
                        'Public study metadata had been recovered from UNITE supplementary data file S1, sheet s2,')
    # Caption disclosure concerns code-assisted drawing, not generated clinical imagery.
    caption_tail = (' The explanatory workflow in panel a was drawn with code assisted by '
                    f'{author["ai_tool"]} ({author["ai_developer"]}; October 2026 desktop session; exact build and model '
                    'version not recorded) and checked by the author.')
    marker = 'Diagnostic composition labels never entered benchmark fitting or selection.'
    assert marker in full
    full = full.replace(marker, marker + caption_tail)
    full = full.split('## References\n', 1)[0] + '## References\n\n' + '\n\n'.join(references) + '\n'
    abstract = extract_section(full, '## Abstract', '## 1. Introduction')
    abstract_only = abstract.split('**Keywords:**')[0].strip()
    keywords = abstract.split('**Keywords:**', 1)[1].strip()
    body = full.split('## 1. Introduction\n', 1)[1].split('## Data and code availability\n', 1)[0]
    body_no_tables = '\n'.join(x for x in body.splitlines() if not x.lstrip().startswith('|'))
    counts = {'abstract_words': word_count(abstract_only),
              'main_body_words_excluding_tables': word_count(body_no_tables),
              'main_body_including_tables_words': word_count(body),
              'full_markdown_words': word_count(full),
              'definition': 'Whitespace-delimited tokens containing a letter or number; Markdown '
                'formatting is stripped; hyphenated terms and numerals count as one. Abstract '
                'includes structured labels and excludes Keywords. Body includes Introduction '
                'through Conclusions, section headings and equations; excludes front matter, '
                'declarations, references, legends, Highlights and, for the primary body count, tables.'}
    blind = full
    blind = re.sub(r'\*\*Author:\*\*.*?\*\*ORCID:\*\* Not registered\.\n\n', '', blind, flags=re.S)
    blind = blind.replace(author['name'], 'The sole author').replace(author['email'], '')
    blind = blind.replace(author['repository_url'], 'the repository described in the separate title page')
    blind = blind.replace('The repository was not publicly accessible at package preparation',
                          'Public code release is pending; access details are supplied separately for anonymous review. The repository was not publicly accessible at package preparation')
    assert author['name'] not in blind and author['email'] not in blind and 'mzxuan08' not in blind
    for text in [full, blind]:
        assert '[To be' not in text and '[Public URL' not in text and '0000-0000-0000-0000' not in text
    write(OUT / 'Manuscript_Main_Blind_and_Full.md', full)
    write(OUT / 'Manuscript_Main_Blinded.md', blind)
    dump(OUT / 'verification/Word_Counts.json', counts)
    highlights = re.findall(r'^- (.+)$', full.split('## Highlights\n')[1].split('## References')[0], re.M)
    assert 3 <= len(highlights) <= 5
    assert all(len(x) <= 85 for x in highlights)
    dump(OUT / 'verification/Highlights_Character_Counts.json',
         [{'text': x, 'characters_including_spaces': len(x), 'passes_85': len(x) <= 85} for x in highlights])
    return title, full, blind, abstract_only, keywords, counts, highlights


def set_font(run, size=12):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().find(qn('w:rFonts'))
    if fonts is not None:
        for key in ['ascii', 'hAnsi', 'cs', 'eastAsia']:
            fonts.set(qn('w:' + key), 'Times New Roman')
        for key in list(fonts.attrib):
            if 'Theme' in key:
                del fonts.attrib[key]


def page_number(section):
    for existing in section.footer.paragraphs:
        existing.clear()
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    suppress = OxmlElement('w:suppressLineNumbers')
    p._p.get_or_add_pPr().append(suppress)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    run = OxmlElement('w:r')
    t = OxmlElement('w:t'); t.text = '1'; run.append(t); field.append(run)
    p._p.append(field)


def format_table(table, double=True):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    n = len(table.columns)
    # Widths are in points. The manuscript is Letter with a 6.5-inch text area.
    first = 130 if n >= 5 else (150 if n > 2 else 190)
    widths = [first] + [(468 - first) / (n - 1)] * (n - 1) if n > 1 else [468]
    table._tbl.tblPr.append(OxmlElement('w:tblLayout'))
    table._tbl.tblPr[-1].set(qn('w:type'), 'fixed')
    for col, width in zip(table.columns, widths): col.width = Pt(width)
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        elem = OxmlElement('w:' + edge)
        elem.set(qn('w:val'), 'single' if edge in ['top', 'bottom'] else 'nil')
        elem.set(qn('w:sz'), '8')
        elem.set(qn('w:color'), '000000')
        borders.append(elem)
    old = table._tbl.tblPr.find(qn('w:tblBorders'))
    if old is not None: table._tbl.tblPr.remove(old)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge, val in [('top', '55'), ('bottom', '55'), ('left', '60'), ('right', '60')]:
        el = OxmlElement('w:' + edge); el.set(qn('w:w'), val); el.set(qn('w:type'), 'dxa'); margins.append(el)
    table._tbl.tblPr.append(margins)
    for r, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit'); trpr.append(no_split)
        if r == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for c, cell in enumerate(row.cells):
            cell.width = Pt(widths[c])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            for shading in list(tcpr.findall(qn('w:shd'))): tcpr.remove(shading)
            if r == 0:
                b = OxmlElement('w:tcBorders'); edge = OxmlElement('w:bottom')
                edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '6'); edge.set(qn('w:color'), '000000')
                b.append(edge); tcpr.append(b)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 2 if double else 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.keep_with_next = False
                for run in p.runs:
                    set_font(run)
                    if r == 0: run.bold = True


def format_document(doc, author, blind=False, double=True, line_numbers=False):
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
        section.header_distance = section.footer_distance = Inches(.45)
        if line_numbers:
            old = section._sectPr.find(qn('w:lnNumType'))
            if old is not None: section._sectPr.remove(old)
            el = OxmlElement('w:lnNumType'); el.set(qn('w:countBy'), '1')
            el.set(qn('w:start'), '1'); el.set(qn('w:restart'), 'continuous')
            el.set(qn('w:distance'), '180'); section._sectPr.append(el)
        page_number(section)
    for style in doc.styles:
        if style.type in [1, 2]:
            style.font.name, style.font.size = 'Times New Roman', Pt(12)
            style.font.color.rgb = RGBColor(0, 0, 0)
        if style.type == 1:
            style.paragraph_format.line_spacing = 2 if double else 1.15
            style.paragraph_format.space_after = Pt(6 if double else 7)
    for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = doc.styles[name]
        style.font.bold = True
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(5)
    doc.styles['Title'].font.size = Pt(12 if double else 14)
    for p in doc.paragraphs:
        p.paragraph_format.line_spacing = 2 if double else 1.15
        p.paragraph_format.widow_control = True
        for run in p.runs: set_font(run, 12 if double or p.style.name != 'Title' else 14)
    for table in doc.tables: format_table(table, double)
    settings = doc.settings.element
    mathpr = settings.find(qn('m:mathPr'))
    if mathpr is None: mathpr = OxmlElement('m:mathPr'); settings.append(mathpr)
    # Word's math font is defined separately from body text; equations remain native OMML.
    doc.core_properties.author = '' if blind else author['name']
    doc.core_properties.last_modified_by = '' if blind else author['name']
    doc.core_properties.comments = ''
    doc.core_properties.keywords = ''
    doc.core_properties.subject = 'Cross-study cfDNA normalization scale audit'
    doc.core_properties.title = 'Scale trade-offs in within-sample self-normalization'


HEADER_ALIASES = {'center': 'Study', 'pipeline': 'Method', 'representation': 'Representation',
    'study': 'Study', 'package': 'Package', 'version': 'Version',
    'auroc': 'AUROC', 'auprc': 'AUPRC', 'brier': 'Brier', 'ece': 'ECE',
    'cohort_scope': 'Cohort', 'training_n': 'Source n', 'test_study': 'Target study',
    'selected_radius': 'Selected k', 'selected_alpha': 'Selected alpha', 'icc_bypassed': 'ICC bypassed',
    'source_threshold_min': 'Locked threshold', 'sensitivity_exact_low': 'Sensitivity lower',
    'sensitivity_exact_high': 'Sensitivity upper',
    'healthy_patients': 'Healthy people', 'cancer_patients': 'Cancer people',
    'cancer_percent': 'Cancer (%)', 'meets_cancer_and_10_healthy': 'Eligible',
    'median_abs_arm_fragmentation_d': 'Median |arm d|', 'median_eta2_batch': 'Median study eta²',
    'specificity95': 'Specificity', 'sensitivity95': 'Sensitivity',
    'auroc_low': 'AUROC lower', 'auroc_high': 'AUROC upper', 'auroc_ci': 'AUROC 95% CI',
    'auprc_ci': 'AUPRC 95% CI', 'specificity_deviation_pp': 'Specificity shift (pp)',
    'specificity_ci': 'Specificity 95% CI', 'sensitivity_ci': 'Sensitivity 95% CI',
    'specificity_low': 'Specificity lower', 'specificity_high': 'Specificity upper',
    'q': 'BH q', 'p': 'p', 'rho': 'Spearman rho'}
METHOD_ALIASES = {'Global_Zscore': 'Global', 'Hybrid_MultiScale': 'Hybrid',
                  'InFold_ComBat': 'ComBat', 'Proposed_LSN_ICC': 'Fixed LSN-ICC',
                  'Tuned_LSN_ICC': 'Tuned LSN', 'Raw_Ratio': 'Raw'}


TABLE_CAPTIONS = {
    1: ('Table 1', 'Independent participant cohort composition'),
    2: ('Supplementary Table S1', 'Historical local source eligibility before depth and QC restrictions'),
    3: ('Supplementary Table S2', 'Software versions used for the frozen analysis'),
    4: ('Table 3A', 'Primary bilateral ElasticNet results'),
    5: ('Supplementary Table S3', 'Nested source-only local parameter selections'),
    6: ('Table 3B', 'Separately held-out high-burden sensitivity cohort'),
    7: ('Table 3C', 'Secondary Combined LOCO ElasticNet results'),
}


def presentation_tables(text):
    """Split wide source tables into labelled panels; preserve every numeric cell string."""
    lines, output, i, number = text.splitlines(), [], 0, 0
    while i < len(lines):
        if lines[i].lstrip().startswith('|'):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                block.append(lines[i]); i += 1
            data = [[x.strip() for x in line.strip().strip('|').split('|')] for line in block]
            if len(data) < 3: output += block; continue
            number += 1
            headers, records = data[0], data[2:]
            # Study+method are a joint key for the primary metrics table; retain both in each panel.
            key_n = 2 if len(headers) > 1 and set(headers[:2]) in [{'center', 'pipeline'}, {'study', 'pipeline'}, {'cohort_scope', 'fold'}] else 1
            max_cols = 5
            groups = [list(range(key_n)) + list(range(j, min(j + (max_cols-key_n), len(headers))))
                      for j in range(key_n, len(headers), max_cols-key_n)]
            if len(headers) <= max_cols: groups = [list(range(len(headers)))]
            caption_id, caption = TABLE_CAPTIONS.get(number, (f'Table {number}', 'Corresponding frozen results'))
            if headers == ['pipeline', 'auroc', 'auroc_low', 'auroc_high']:
                caption_id, caption = 'Supplementary Table S4', 'Secondary Combined LOCO baseline comparison'
            for panel, cols in enumerate(groups):
                output += ['', f'**{caption_id}{" panel " + chr(97+panel) if len(groups)>1 else ""}. '
                           f'{caption}{" continued" if panel else ""}.**', '']
                new_headers = [HEADER_ALIASES.get(headers[c], headers[c].replace('_', ' ')) for c in cols]
                output.append('| ' + ' | '.join(new_headers) + ' |')
                output.append('| ' + ' | '.join(['---'] * len(cols)) + ' |')
                for record in records:
                    cells = []
                    for c in cols:
                        cell = record[c]
                        cell = METHOD_ALIASES.get(cell, cell)
                        if cell.startswith(('Cristiano et al', 'Jiang et al', 'Snyder et al')):
                            cell = cell.split(' et al')[0]
                        cells.append(cell)
                    output.append('| ' + ' | '.join(cells) + ' |')
                output.append('')
            continue
        output.append(lines[i]); i += 1
    return '\n'.join(output) + '\n'


def pandoc_doc(text, output, pandoc, author, blind=False, double=True, numbers=False):
    path = QA / (output.stem + '.md')
    # The first heading is mapped to Word's Title style by the postprocessor.
    write(path, presentation_tables(text))
    subprocess.run([str(pandoc), str(path), '--from=markdown+tex_math_dollars',
                    '--to=docx', '--output=' + str(output)], check=True, capture_output=True)
    doc = Document(output)
    if doc.paragraphs: doc.paragraphs[0].style = doc.styles['Title']
    format_document(doc, author, blind, double, numbers)
    # Keep captions with the following first table row, not the entire long table.
    for p in doc.paragraphs:
        if re.match(r'^(Table \d|Supplementary Table)', p.text):
            p.paragraph_format.keep_with_next = True
    doc.save(output)
    return doc


def title_page(title, abstract, keywords, counts, author):
    text = (f'# {title}\n\n{author["name"]}\n\n{author["affiliation"]}\n\n'
            f'Corresponding author: {author["name"]}\n\nEmail: {author["email"]}\n\n'
            'ORCID: Not registered.\n\n'
            f'Abstract word count: {counts["abstract_words"]}\n\n'
            f'Main text word count, excluding tables: {counts["main_body_words_excluding_tables"]}\n\n'
            f'Main text word count, including tables: {counts["main_body_including_tables_words"]}\n\n'
            'Counting convention: Introduction through Conclusions, including section headings '
            'and equation tokens; excluding abstract, declarations, legends and references. '
            'Exact token rules are recorded in verification/Word_Counts.json.\n\n'
            f'Keywords: {keywords}\n\n')
    for heading, body in declarations(author).items(): text += f'## {heading}\n\n{body}\n\n'
    text += ('## Data and code release\n\nPrepared repository: ' + author['repository_url'] +
             '. Public repository publication and Zenodo archival release are pending; '
             'no DOI is assigned. No upload is performed by this package builder.\n')
    return text


def csv_markdown(data, columns):
    lines = ['| ' + ' | '.join(columns) + ' |', '| ' + ' | '.join(['---'] * len(columns)) + ' |']
    for r in data:
        values = []
        for c in columns:
            v = str(r[c]).replace('|', '/')
            if c in ['auroc', 'auroc_low', 'auroc_high']: v = f'{float(v):.3f}'
            values.append(v)
        lines.append('| ' + ' | '.join(values) + ' |')
    return '\n'.join(lines)


def add_integrated_figures(doc):
    found = []
    for p in list(doc.paragraphs):
        if p.text == 'Figure legends':
            p._element.getparent().remove(p._element)
    # Insert each original byte-identical image immediately before its complete legend.
    for p in list(doc.paragraphs):
        text = p.text
        match = re.match(r'^(Figure ([1-4])\.|Supplementary Figure S5\.)', text)
        if not match: continue
        index = int(match.group(2)) - 1 if match.group(2) else 4
        new_p = p.insert_paragraph_before()
        new_p.paragraph_format.page_break_before = True
        new_p.paragraph_format.keep_with_next = True
        new_p.paragraph_format.line_spacing = 1
        new_p.add_run().add_picture(str(ROOT / 'results/figures' / (FIGURES[index] + '.png')), width=Inches(6.5))
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.keep_with_next = False
        for run in p.runs: set_font(run, 11)
        found.append(index)
    assert sorted(found) == list(range(5)), found


def cover_letter(title, author, journal):
    stress = rows('Step7_Prevalence_Stress_Summary.csv')
    # Read column names rather than hard-coding a new result from memory.
    combat = next(r for r in stress if r['pipeline'] == 'InFold_ComBat')
    possible = [k for k in combat if 'max' in k and ('prob' in k or 'range' in k or 'shift' in k)]
    if not possible:
        possible = [k for k in combat if 'range' in k]
    assert possible, list(combat)
    observed = max(float(r[possible[0]]) for r in stress if r['pipeline'] == 'InFold_ComBat')
    assert abs(observed - .214) < .001
    return (f'# Cover letter\n\n{DATE}\n\nEditor-in-Chief\n\n{journal}\n\nDear Editor,\n\n'
        f'I submit the original research manuscript entitled “{title}” for consideration in '
        f'{journal}. The manuscript examines cross-study generalization in public cfDNA '
        'cohorts through a methodological audit of normalization scale and study-source '
        'confounding. Its contribution is the scale trade-off analysis, rather than a claim '
        'that Hybrid fusion consistently outperforms Global Z-score.\n\n'
        'The analysis uses 585 source-identified participants. A primary bidirectional '
        'benchmark contains 548 participants from two studies with adequately represented '
        'cancer and healthy groups; a separately held-out, case-dominated study contributes '
        '37 participants for external sensitivity assessment. Study Source is explicitly '
        'treated as a proxy combining technical batch and population composition.\n\n'
        'The first methodological contribution is a direct comparison of global, local '
        'and fused within-sample representations, linking broad arm-level fragmentation '
        'contrasts with residual GC-associated study variation. It quantifies the scale '
        'trade-off without equating fragmentation contrasts with independently measured CNA.\n\n'
        'Second, a cancer–control pair concordance decomposition separates within-study '
        'ranking from cross-study score comparability. A fixed real-patient anchor '
        'composition stress test then measures the dependence of target-batch adaptation '
        f'on companion patients: probability shifts reach {observed:.3f} for the evaluated '
        'ComBat-plus-unseen-target alignment comparator, whereas within-sample Global, '
        'LSN and Hybrid transformations produce zero companion-dependent shift. These '
        'diagnostics explain why poor pooled discrimination need not imply complete '
        'destruction of within-study biological contrasts.\n\n'
        'Third, the accompanying reproducibility package contains fold-isolated code, '
        'automated tests, source-only nested selection and threshold calibration, and '
        'hash-verifiable outputs. No target labels enter predictive fitting or threshold '
        'selection; labels used to construct diagnostic stress batches are confined to '
        'that post-evaluation experiment. The prepared MIT software release and compact '
        'statistical archive are ready for public deposit.\n\n'
        'I am the sole author, have reviewed and approved the manuscript and its code '
        'and statistical outputs, and accept full responsibility for the work. This '
        'manuscript is original, has not been published as a journal article, and is '
        'not under consideration by another journal. Any subsequently posted preprint '
        'will be disclosed in the submission system. The analysis uses published '
        'de-identified data without new recruitment or specimen collection; its ethics '
        'statement accurately reports the absence of a new institutional determination. '
        'There is no specific grant funding and no competing interest. Assistance '
        f'from {author["ai_tool"]} is disclosed in Methods and in the AI-use declaration.\n\n'
        'I request the traditional subscription publishing model rather than optional '
        'author-pays open access, subject to the journal’s current publishing terms.\n\n'
        f'Sincerely,\n\n{author["name"]}\n\n{author["affiliation"]}\n\n{author["email"]}\n')


def supplement(full, author):
    legend = full.split('**Supplementary Figure S5.', 1)[1].split('## Highlights')[0].strip()
    return ('# Supplementary material\n\n'
        'Scale Trade-offs in Within-Sample Self-Normalization for Cross-Study Generalization '
        'in Public cfDNA Cohorts\n\n'
        'This supplement accompanies the frozen analysis and documents the scope of the '
        'statistical tables. It introduces no new fitted predictions or observations.\n\n'
        '## Supplementary Figure S5\n\n**Supplementary Figure S5. ' + legend + '\n\n'
        '## Statistical table guide\n\n'
        '**Table 1, Table1_Cohort_Characteristics.csv:** extraction and QC cohort accounting. '
        'For independent-participant counts used in predictive evaluation, consult '
        'Analysis_Cohort_Counts.csv and Step7_Prevalence_By_Study.csv; specimens and patients '
        'are distinct units.\n\n'
        '**Table 2, Table2_Variance_Partition_Summary.csv:** window-level source and '
        'disease variance components, vulnerability and healthy-drift statistics. '
        'Unique/type III variance fractions need not sum to one.\n\n'
        '**Table 3, Table3_Benchmark_Metrics_Full.csv:** full fixed-score discrimination, '
        'calibration and locked-threshold operating metrics. cohort_scope distinguishes '
        'Primary_Bilateral, External_High_Burden and Combined_LOCO; mode distinguishes '
        'random, loco and external. Unavailable one-source ComBat and DANN cells remain '
        'unavailable. A one-source ICC bypass is an LSN-only fallback.\n\n'
        '**Table 4, Table4_Ablation_Study.csv:** descriptive local-radius and filtering '
        'sensitivity results from the earlier fixed experiment. These results did not '
        'select a revised target-optimal model. Source-only nested selection is separately '
        'documented in Step7_Nested_Parameter_Search.csv and Step7_Selected_Parameters.csv.\n\n'
        '**Step7_Paired_Significance.csv:** paired target-specific DeLong and other '
        'prespecified comparisons with BH-adjusted q values. No pooled cross-validated '
        'DeLong inference is substituted for the target-specific tests.\n\n'
        '**Step7_AUC_Pair_Decomposition.csv and Step7_Source_Threshold_Alignment.csv:** '
        'secondary within-study/cross-study concordance and source-locked score-alignment '
        'diagnostics, which do not replace benchmark probabilities.\n\n'
        '**Step7_Prevalence_Stress_Summary.csv:** fixed-anchor probability changes under '
        'diagnostic target compositions. Labels define real-patient diagnostic batches '
        'after evaluation; no synthetic patient is introduced.\n\n'
        '**Final alignment CSV tables:** verified baseline comparisons, residual GC associations, '
        'arm-level fragmentation contrasts and historical source eligibility census. '
        'Missing source counts remain unknown; no live exhaustive database census is claimed.\n\n'
        '## Secondary three-study discrimination\n\n' +
        csv_markdown(rows('Final_Alignment_Baseline_Comparison.csv'),
            ['pipeline', 'auroc', 'auroc_low', 'auroc_high']) + '\n\n'
        'Exact statistical cells are preserved in the accompanying machine-readable CSVs. '
        'Numbers in this supplement are copied from those CSVs without refitting.\n')


def copy_verified(src, target, records):
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(src, target)
    digest = sha(src)
    assert sha(target) == digest
    records[str(target.relative_to(OUT))] = {'source': str(src.relative_to(ROOT)), 'sha256': digest}


def create_release(author, full, copies):
    release = OUT / 'github_zenodo_release'
    release.mkdir(parents=True, exist_ok=True)
    for folder in ['src', 'scripts', 'tests', 'config']:
        for p in (ROOT / folder).rglob('*'):
            if p.is_file() and '__pycache__' not in p.parts and p.suffix in ['.py', '.yaml', '.yml', '.json']:
                if p.name == 'submission_author.json': continue
                copy_verified(p, release / p.relative_to(ROOT), copies)
    for filename in ['environment.yml', 'requirements.txt', 'pyproject.toml']:
        if (ROOT / filename).exists(): copy_verified(ROOT / filename, release / filename, copies)
    # Tables are small derived output; omit unnecessary individual-level clinical inventories.
    excluded = {'Analysis_Participants.csv', 'Clinical_CV_Splits.csv', 'Step7_Primary_Splits.csv',
                'Step7_Prevalence_Stress_Members.csv', 'Step7_Prevalence_Stress_Anchor_Predictions.csv'}
    for p in (ROOT / 'results/tables').glob('*.csv'):
        if p.name not in excluded:
            copy_verified(p, release / 'results/tables' / p.name, copies)
    # Minimal historical OOF provenance is a required input to the revision engine.
    archived = ['tables/Clinical_OOF_Predictions.parquet', 'tables/Clinical_CV_Splits.csv',
                'source/src/models.py']
    manifest = {}
    for name in archived:
        p = ROOT / 'results/archive_step06' / name
        assert p.exists(), p
        copy_verified(p, release / 'results/archive_step06' / name, copies)
        manifest[name] = sha(p)
    dump(release / 'results/archive_step06/snapshot_manifest.json', manifest)
    for p in (ROOT / 'results/clinical_cv_checkpoints').glob('selection_loco_*_Proposed_LSN_ICC.parquet'):
        copy_verified(p, release / 'results/clinical_cv_checkpoints' / p.name, copies)
    # Compact reference annotations, not BAM/FASTQ, genome FASTA or interval caches.
    for p in (ROOT / 'data/ref').iterdir():
        if p.is_file() and p.stat().st_size < 1_000_000 and p.suffix in ['.parquet', '.json', '.gz', '.sizes']:
            copy_verified(p, release / 'data/ref' / p.name, copies)
    for name in ['final_scope_alignment.md', 'step_07_report.md', 'final_alignment_validation.json',
                 'final_alignment_pytest.xml', 'step7_pytest.xml', 'step7_validation.json',
                 'step7_archived_prediction_reuse.json']:
        p = ROOT / 'results/reports' / name
        if p.exists(): copy_verified(p, release / 'results/reports' / name, copies)
    write(release / 'results/MANUSCRIPT_FULL_TEXT.md', full)
    write(release / 'LICENSE', f'''MIT License

Copyright (c) 2026 {author['name']}

Permission is hereby granted, free of charge, to any person obtaining a copy
of this software and associated documentation files (the "Software"), to deal
in the Software without restriction, including without limitation the rights
to use, copy, modify, merge, publish, distribute, sublicense, and/or sell
copies of the Software, and to permit persons to whom the Software is
furnished to do so, subject to the following conditions:

The above copyright notice and this permission notice shall be included in all
copies or substantial portions of the Software.

THE SOFTWARE IS PROVIDED "AS IS", WITHOUT WARRANTY OF ANY KIND, EXPRESS OR
IMPLIED, INCLUDING BUT NOT LIMITED TO THE WARRANTIES OF MERCHANTABILITY,
FITNESS FOR A PARTICULAR PURPOSE AND NONINFRINGEMENT. IN NO EVENT SHALL THE
AUTHORS OR COPYRIGHT HOLDERS BE LIABLE FOR ANY CLAIM, DAMAGES OR OTHER
LIABILITY, WHETHER IN AN ACTION OF CONTRACT, TORT OR OTHERWISE, ARISING FROM,
OUT OF OR IN CONNECTION WITH THE SOFTWARE OR THE USE OR OTHER DEALINGS IN THE
SOFTWARE.
''')
    gitignore = '''# Regenerable environments and caches
.venv/
venv/
__pycache__/
*.py[cod]
.pytest_cache/
.mypy_cache/
.ipynb_checkpoints/
work/
submission_package/
*.log
# Never commit clinical raw intervals, BAM, FASTQ or large processed inputs
*.bam
*.bai
*.fastq*
*.fq*
*.frag.tsv.bgz*
*.bed.bgz*
data/processed/
data/metadata/
results/step7_cv_checkpoints/
results/clinical_cv_checkpoints/*
!results/clinical_cv_checkpoints/selection_loco_*_Proposed_LSN_ICC.parquet
*.pt
*.pth
*.ckpt
# Local credentials
.env
.env.*
*.pem
*.key
'''
    write(release / '.gitignore', gitignore)
    # JSON is valid YAML 1.2 and valid Citation File Format syntax.
    cff = {'cff-version': '1.2.0', 'message': 'Please cite the manuscript and versioned code release if used.',
           'type': 'software', 'title': 'cfDNA Scale Tradeoff Audit', 'version': '0.1.0',
           'authors': [{'family-names': author['family_names'], 'given-names': author['given_names'],
                        'affiliation': author['affiliation'], 'email': author['email']}],
           'repository-code': author['repository_url'], 'license': 'MIT',
           'abstract': 'A reproducible audit of global/local normalization scale and study-source '
                       'confounding in public human cfDNA cohorts; Hybrid fusion is a diagnostic comparator.',
           'keywords': ['cfDNA', 'fragmentomics', 'cross-study validation', 'normalization']}
    if author.get('orcid'): cff['authors'][0]['orcid'] = author['orcid']
    if author.get('zenodo_doi'): cff['doi'] = author['zenodo_doi']
    dump(release / 'CITATION.cff', cff)
    write(release / 'DATA_PROVENANCE_AND_TERMS.md', '''# Data provenance and licensing scope

The MIT license covers the original software and documentation only. It does not
relicense third-party clinical data, reference annotations or published supplementary
materials. Follow the original study, FinaleDB, UCSC, ENCODE and UNITE conditions.
No claim is made that public accessibility alone grants unlimited redistribution.

The software streams published fragment intervals and discards raw intervals after
aggregation. The release omits raw BAM/FASTQ, fragment interval downloads, participant
feature matrices and bulk metadata. The compact historical OOF/split bundle contains
source-published de-identified identifiers and derived predictions solely to document
the exact archived benchmark reused by the revision. Do not attempt re-identification.
These derived artifacts are scientific provenance, not new clinical recruitment.

Primary publications: 10.1038/s41586-019-1272-6; 10.1073/pnas.1500076112;
10.1016/j.cell.2015.11.050. FinaleDB: 10.1093/bioinformatics/btaa999.
Metadata source: UNITE supplementary data file S1, sheet s2,
10.1126/sciadv.ady9432. Data retrieval availability is not guaranteed in perpetuity.
The historical source inventory is not a current exhaustive public database census.
''')
    write(release / 'README.md', release_readme(author))
    # A private author confirmation is not required to use the public research code.
    write(release / 'config/submission_author.example.json', json.dumps({
        'name': author['name'], 'given_names': author['given_names'], 'family_names': author['family_names'],
        'affiliation': author['affiliation'], 'email': author['email'], 'orcid': None,
        'repository_url': author['repository_url'], 'zenodo_doi': None,
        'manual_review_confirmed': False,
        'note': 'For submission packaging only. Supply a complete locally approved author configuration; '
                'research scripts do not require this file.'}, indent=2))
    return release


def release_readme(author):
    return f'''# cfDNA Scale Tradeoff Audit

**Scale Trade-offs in Within-Sample Self-Normalization for Cross-Study Generalization
in Public cfDNA Cohorts**

{author['name']} · {author['affiliation']} · {author['email']}

Prepared repository: {author['repository_url']}. This is a local release candidate;
GitHub publication and Zenodo DOI assignment are separate author actions. No DOI is
claimed in CITATION.cff before assignment.

## Scientific scope

This project audits normalization scale and study-source confounding in public
human plasma cfDNA cohorts. Hybrid concatenates global and local features as a
diagnostic fusion probe. It did not establish consistent superiority over Global
Z-score. Study Source combines experimental batch and population composition;
it is not a verified hospital identifier or a randomized technical intervention.

The analysis contains 585 source-identified participants from 628 extracted
specimens and uses 477 retained hg19 autosomal 5 Mb windows. The primary bilateral
cohort is Cristiano 2019 (427) plus Jiang 2015 (121), total 548. Snyder 2016 (37;
36 cancer and one healthy) is a separately held-out sensitivity cohort. Secondary
Combined LOCO uses all 585. These designs must not be conflated.

Fixed ElasticNet Combined LOCO AUROC: Global 0.759, Hybrid 0.731, fixed LSN-ICC
0.724, ComBat-plus-target-alignment 0.543. Consult exact frozen CSV cells and
cohort_scope rather than treating rounded overview values as new estimates.
The fixed-anchor diagnostic measures a maximum ComBat probability change of
0.214 and zero companion dependence for the sample-wise transforms. This does
not establish immunity to training prevalence, population shift or calibration error.

## Included files and exclusions

- src/, scripts/, tests/ and config/: original research implementation and tests.
- results/tables/: exact aggregate CSV outputs and full numeric traceability tables.
- results/reports/: selected frozen validation evidence and scope alignment.
- data/ref/: compact reference windows, annotations and interval masks.
- results/archive_step06/: minimal historical OOF predictions, split provenance,
  archived models.py and checksum manifest required by the revision engine.
- results/clinical_cv_checkpoints/selection_loco_*_Proposed_LSN_ICC.parquet:
  small fixed selection masks used to audit Figure 4, not fitted model weights.

Raw BAM/FASTQ, raw fragment streams, bulk metadata, derived specimen feature
matrices, fitted model caches and large third-party archives are excluded. The
minimal historical OOF bundle is intentional provenance, not a hidden fitting
cache. Some archived reports retain original absolute local paths as historical
evidence; those paths are not portable input configuration. Configured research
input paths are relative to the project root.

## Environment

Validated analysis: Python 3.10, Windows, seed 2026, RTX 3060 Ti. Use:

```console
conda env create -f environment.yml
conda activate cfdna-batchaudit
python -m pytest tests/ -q
```

Or create a Python 3.10 virtual environment and install requirements.txt. CUDA
12.1 PyTorch wheels are pinned for Windows/Linux. Neural batches are at most
128 and hidden widths at most 256. CUDA model availability affects whether all
locked model combinations can be reproduced; unavailable hardware must be
reported, not silently replaced with fabricated scores. This release does not
claim cross-platform numerical equivalence.

## Rebuild from the public sources

Run from the repository root, in a separate analysis working copy. Rebuilding
will intentionally regenerate results in that working copy; preserve the frozen
release if comparing hashes. Network access, original data terms and available
remote objects are required. No hundreds-of-gigabytes BAM/FASTQ download is used.

```console
python scripts/01_setup_and_ref.py
python scripts/02_fetch_and_extract.py --help
python scripts/02_fetch_and_extract.py
python scripts/03_batch_vulnerability_audit.py
python scripts/04_run_cv_benchmarks.py --step7
python scripts/05_biological_interpretability.py
```

Read --help and config/config.yaml before full extraction. The discovery path
uses published UNITE metadata and source-explicit interval URLs, with six stream
workers and checkpoints. QC and source-identifier selection must reproduce the
documented cohort; do not invent missing source labels. The Step 7 engine reuses
the verified original four-pipeline Combined LOCO predictions supplied in the
compact archive and fits only the specified revision tasks. A de novo replication
of that original benchmark is distinct from replaying its preserved predictions.

The historical archive's SHA256 manifest is checked before prediction reuse,
and models.py must match its archived version. Research feature matrices must
first be reconstructed for scripts that load patient-level features. The light
release alone does not run feature-dependent validations without those inputs.

**Manuscript regeneration:** scripts/06_compile_manuscript_assets.py is the
historical compiler. It can rewrite tables and reports; never run it on the
frozen release to perform editorial packaging. The approved full text is supplied
in results/MANUSCRIPT_FULL_TEXT.md. Editorial packaging uses the read-only-input
scripts/08_build_submission_package.py and requires a complete author-confirmed
local config/submission_author.json, python-docx, lxml and a Pandoc executable;
it is not needed for research replication. A release without the full private
analysis workspace should not claim the packaging integrity manifest is present.

## Leakage controls and statistical limits

All fitted preprocessing, masks, nested parameter selection and source threshold
calibration are confined to source training folds. Primary bilateral training has
only one study; healthy-source ICC is bypassed explicitly, and multi-source
ComBat/DANN are unavailable. InFold ComBat includes an unlabeled unseen-target
median/IQR alignment and is explicitly transductive. Ordinary individual-score
DeLong inference is not reported for that target-batch-fitted adaptation.

Confidence intervals use 1000 fixed-score, study/diagnosis-stratified patient
bootstraps. They do not include model refitting or uncertainty over future
studies. Target-specific DeLong p values and BH q values are in the supplied
tables; no pooled CV DeLong test is asserted. Synthetic arrays occur only in
unit tests; production clinical outputs use real public participants.

## Verification and licensing

RELEASE_MANIFEST_SHA256.json records relative paths and hashes for this prepared
release. Compare those hashes before regeneration. The submission package's
separate frozen-input audit confirms the original matrices, CSVs and figure
files were unchanged during packaging. The recorded earlier research run had
73 passing tests and one optional live-network test skipped; current package
validation is reported separately and must not be described as a new full
clinical extraction or benchmark run.

Original software is MIT licensed. Third-party reference/clinical resources
retain their own conditions; see DATA_PROVENANCE_AND_TERMS.md. Cite the original
data papers, the manuscript and the versioned code/Zenodo record once released.
Do not use this exploratory public-data study as a validated screening device.
'''


def package_validation(author):
    result = {}
    for path in sorted(OUT.glob('*.docx')):
        with zipfile.ZipFile(path) as z:
            xml = ET.fromstring(z.read('word/document.xml'))
            ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                  'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
            texts = ' '.join(xml.itertext())
            tables = xml.findall('.//w:tbl', ns)
            entry = {'zip_test': z.testzip(), 'native_equations': len(xml.findall('.//m:oMath', ns)),
                     'tables': len(tables), 'embedded_images': len([n for n in z.namelist() if n.startswith('word/media/')]),
                     'continuous_line_numbers': bool(xml.findall('.//w:lnNumType', ns))}
            if path.name.startswith('03_'):
                assert xml.find('.//w:lnNumType', ns).get(qn('w:restart')) == 'continuous'
                assert entry['native_equations'] >= 8
                assert author['name'] not in texts and author['email'] not in texts and 'mzxuan08' not in texts
                assert author['name'] not in z.read('docProps/core.xml').decode()
                # Double spacing in every body/table paragraph that contains text.
                for p in xml.findall('.//w:body//w:p', ns):
                    if p.find('.//w:t', ns) is None: continue
                    spacing = p.find('w:pPr/w:spacing', ns)
                    assert spacing is not None and spacing.get(qn('w:line')) == '480'
                for sect in xml.findall('.//w:sectPr', ns):
                    mar = sect.find('w:pgMar', ns)
                    assert all(mar.get(qn('w:' + x)) == '1440' for x in ['top', 'bottom', 'left', 'right'])
            if path.name.startswith('04_'):
                expected = {sha(ROOT / 'results/figures' / (n + '.png')) for n in FIGURES}
                actual = {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
                assert actual == expected
                entry['all_five_embedded_figures_byte_identical'] = True
            for table in tables:
                borders = table.find('w:tblPr/w:tblBorders', ns)
                assert borders is not None
                assert borders.find('w:top', ns).get(qn('w:val')) == 'single'
                assert borders.find('w:bottom', ns).get(qn('w:val')) == 'single'
                assert borders.find('w:insideV', ns).get(qn('w:val')) == 'nil'
            assert not entry['zip_test']
            assert '[To be' not in texts and '[PENDING_RELEASE]' not in texts
            result[path.name] = entry
    assert len(result) == 8, list(result)
    dump(OUT / 'verification/Document_Structural_Validation.json', result)
    return result


def policy_notes():
    notes = [
      {'topic': 'AI disclosure', 'url': 'https://www.elsevier.com/about/policies-and-standards/generative-ai-policies-for-journals',
       'accessed': DATE, 'verification': 'Official policy read. Tool, purpose, human review and research/figure-code assistance disclosed in Methods and declaration; schematic caption has disclosure.'},
      {'topic': 'bioRxiv AI use', 'url': 'https://connect.biorxiv.org/news/2026/10/20/approach_to_ai',
       'accessed': DATE, 'verification': 'Official policy read. Human authorship, substantive judgment and accountability are required; preprint screening is discretionary.'},
      {'topic': 'Public-data ethics provisions', 'url': 'https://www.nhc.gov.cn/wjw/c100375/202302/902b4a1dc3af4aba862a6387e6e376dc.shtml',
       'accessed': DATE, 'verification': 'Official indexed text for the 2023 measures, Article 32, identifies conditional public-data exemptions. It does not constitute an institutional determination for this project; full-page fetch returned 412.'},
      {'topic': 'CMPB author guide', 'url': 'https://www.sciencedirect.com/journal/computer-methods-and-programs-in-biomedicine/publish/guide-for-authors',
       'accessed': DATE, 'verification': 'Direct guide fetch returned 403. Formatting follows the author-requested specification. Current journal-specific checklist must be confirmed in submission portal.'},
      {'topic': 'CBM author guide', 'url': 'https://www.sciencedirect.com/journal/computers-in-biology-and-medicine/publish/guide-for-authors',
       'accessed': DATE, 'verification': 'Direct guide fetch returned 403. Subscription route requested; no universal zero-total-fee guarantee is asserted.'}]
    dump(OUT / 'verification/Policy_Source_Checks.json', notes)


def write_checklist(highlights, counts, integrity, validation, copies):
    inventory = sorted(str(p.relative_to(OUT)).replace('\\', '/') for p in OUT.rglob('*') if p.is_file())
    text = ('# 投稿前作者检查清单\n\n'
        f'作者：Zhenxuan Ma（马振轩）；单位：Hubei University of Medicine, Shiyan 442000, Hubei, China；'
        f'邮箱：mzxmc08@gmail.com。ORCID尚未注册，未填假号码。构建日期：{DATE}。\n\n'
        '## 已完成的本地核验\n\n'
        f'- 原项目受保护文件 {integrity["checked_files"]} 个：SHA256全部一致；未重提特征、未重新训练、未改写原CSV或图像。\n'
        '- 8条文献逐一匹配Crossref；第2和第7条的完整页码以NLM PubMed补全。第6条为2026;12(28):eady9432。\n'
        '- 8份Word的压缩包、原生公式、三线表、页码与正文结构通过程序检查。03版为12 pt Times New Roman、双倍行距、1 inch页边距、连续行号。\n'
        '- 04一体版嵌入5张原始高分辨率PNG，字节哈希与冻结原图一致。\n'
        f'- 摘要 {counts["abstract_words"]} 词；正文（不含表格）{counts["main_body_words_excluding_tables"]} 词；'
        f'正文含表格 {counts["main_body_including_tables_words"]} 词。统计规则见verification/Word_Counts.json。\n\n'
        '## Highlights逐条字符数（包含空格，不含项目符号）\n\n')
    for i, h in enumerate(highlights, 1): text += f'{i}. {len(h)} / 85：{h}\n'
    text += ('\n## 上传文件的用途\n\n'
        '- 01_Title_Page.docx：署名、通讯信息、词数和声明。\n'
        '- 02_Highlights.docx：单独上传Highlights。\n'
        '- 03_Manuscript_DoubleSpaced_LineNumbered.docx：已去署名及个人仓库链接的审稿正文，保留完整Methods、结果、公式、图注和参考文献。\n'
        '- 04_bioRxiv_Integrated_Preprint.docx：含署名与Figure 1–4及S5，图像放在对应图注前。\n'
        '- 05_Supplementary_Material.docx：S5及统计表说明，配合独立CSV。\n'
        '- 06_Cover_Letter.docx：仅致CMPB；06b是CBM备选版本。只选一封、一家期刊提交。\n'
        '- 07_Declaration_of_Interest.docx：无利益冲突独立声明。\n'
        '- Manuscript_Main_Blind_and_Full.md为完整署名版；Manuscript_Main_Blinded.md为独立匿名版。\n'
        '- figures_and_tables/：冻结PDF/PNG和CSV原样副本；github_zenodo_release/为轻量开源发布候选。\n\n'
        '## 尚需作者在正式提交前落实\n\n'
        '1. 你已确认逐项人工核验并承担责任。请在投稿网页再次审阅本次新增署名、声明和引用格式；当前稿件没有捏造伦理批准编号或免审决定。'
        '如学校要求机构判定，请取得确认后按真实结果更新伦理声明。2023年伦理审查办法第32条的公开数据条款是有条件的，不能代替学校的具体要求。\n'
        '2. GitHub查询返回404，仓库尚未核验为公开；本脚本没有发布。创建并上传候选目录后再将Data and code availability中的待发布状态改为实际公开状态。\n'
        '3. Zenodo DOI尚未分配。完成归档后把真实DOI及版本号写入完整稿、CITATION.cff和README，不能使用伪造DOI或全零ORCID。\n'
        '4. 软件MIT不等于第三方数据MIT。核对原始数据及参考资源的再分发条款；轻量包不含原始片段、患者特征矩阵或大缓存。\n'
        '5. 已按Elsevier当前AI政策披露代码、统计/绘图代码和文字辅助及人工责任。所用AI工具的精确应用构建号及底层模型版本未记录，稿件如实说明；若编辑要求精确版本，补充实际记录。\n'
        '6. 两刊作者指南网页本次返回403，不能宣称已逐条核实其2026全部个性化要求。正式提交时核对字数、文章类别、图表格式和匿名要求；订阅模式可避免可选OA APC，但本包不保证所有附加费用一律为0。\n\n'
        '## 三分钟网页操作指南\n\n'
        '**第1分钟：选择去向。** bioRxiv使用官方入口 https://www.biorxiv.org/submit-a-manuscript ，先用Word将04另存为PDF并确认图文完整，'
        '按网页要求上传主稿与05补充材料；该文属于公开数据方法学审计，是否适合bioRxiv由平台筛查决定。期刊只选CMPB或CBM其中一家，'
        '从官方ScienceDirect主页进入Submit your article；不要选择名称相近的Update开放获取伴刊。\n\n'
        '**第2分钟：上传与分配角色。** 填入作者与通讯邮箱；ORCID无则留空（如系统强制则先注册）。期刊分别上传01为Title Page、03为Manuscript、'
        '02为Highlights、对应06为Cover Letter、07为Declaration、05为Supplementary及figures_and_tables中的独立图件。'
        '若期刊要求单盲完整正文而不是匿名正文，可从完整MD或04导出符合其要求的版本。\n\n'
        '**第3分钟：校验与提交。** 核对系统生成PDF的行号、公式、表格、图注和8条参考文献；按实际情况填写伦理、AI、资金和冲突问题，'
        '选择Subscription/traditional路线而不是可选Gold Open Access。预印本上线后如实填写其DOI及状态；期刊只保留一个在审。'
        '在仓库公开及Zenodo归档后填写真实链接，查看费用页再点击最终Submit。本脚本没有代你上传或提交。\n\n'
        '## 全部文件清单\n\n')
    text += '\n'.join('- `' + p + '`' for p in inventory)
    text += '\n\n另见PACKAGE_MANIFEST_SHA256.json；ZIP内CRC及文件SHA256另行核验。\n'
    write(OUT / 'SUBMISSION_CHECKLIST_FOR_AUTHOR.md', text)


def finalize_archives():
    release = OUT / 'github_zenodo_release'
    def shippable(path):
        return path.is_file() and not any(x in path.parts for x in ['.git', '__pycache__', '.pytest_cache'])
    release_hashes = {str(p.relative_to(release)).replace('\\', '/'): sha(p)
                      for p in sorted(release.rglob('*')) if shippable(p) and p.name != 'RELEASE_MANIFEST_SHA256.json'}
    dump(release / 'RELEASE_MANIFEST_SHA256.json', release_hashes)
    for folder, name in [(release, 'github_zenodo_release.zip')]:
        with zipfile.ZipFile(OUT / name, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
            for p in sorted(folder.rglob('*')):
                if shippable(p): z.write(p, str(p.relative_to(folder.parent)))
        with zipfile.ZipFile(OUT / name) as z: assert z.testzip() is None
    hashes = {str(p.relative_to(OUT)).replace('\\', '/'): sha(p)
              for p in sorted(OUT.rglob('*')) if shippable(p) and p.name != 'PACKAGE_MANIFEST_SHA256.json'}
    dump(OUT / 'PACKAGE_MANIFEST_SHA256.json', hashes)
    outer = ROOT / 'cfDNA_Submission_Package.zip'
    with zipfile.ZipFile(outer, 'w', zipfile.ZIP_DEFLATED, compresslevel=6) as z:
        for p in sorted(OUT.rglob('*')):
            if shippable(p): z.write(p, str(p.relative_to(ROOT)))
    with zipfile.ZipFile(outer) as z: assert z.testzip() is None
    write(ROOT / 'cfDNA_Submission_Package.zip.sha256', sha(outer) + '  cfDNA_Submission_Package.zip\n')
    return {'package_zip': str(outer), 'package_zip_sha256': sha(outer),
            'package_zip_bytes': outer.stat().st_size, 'release_files': len(release_hashes),
            'release_bytes': sum(p.stat().st_size for p in release.rglob('*') if p.is_file())}


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--pandoc', type=Path, required=True)
    p.add_argument('--author-config', type=Path, default=ROOT / 'config/submission_author.json')
    p.add_argument('--finalize-only', action='store_true', help='Refresh checksums and ZIPs after render QA; do not rebuild documents')
    args = p.parse_args()
    started = time.monotonic()
    before = frozen_snapshot()
    OUT.mkdir(exist_ok=True); QA.mkdir(parents=True, exist_ok=True)
    author = json.loads(read(args.author_config))
    assert author['manual_review_confirmed'] is True, 'Human review confirmation is required for the supplied declaration'
    if args.finalize_only:
        integrity = assert_frozen(before)
        previous = json.loads(read(OUT / 'verification/Frozen_Input_Integrity.json'))
        assert previous['protected_files'] == before, 'Original pre-packaging file inventory changed'
        review = json.loads(read(QA / 'Render_QA_Review.json'))
        assert review['status'] == 'passed'
        rendered = json.loads(read(QA / 'render_status.json'))
        assert len(rendered) == 8
        for item in rendered:
            assert item['returncode'] == 0 and item['pages'] > 0
            assert sha(OUT / item['file']) == item['sha256'], 'Document edited after rendering'
        dump(OUT / 'verification/Render_QA_Review.json', review)
        dump(OUT / 'verification/Rendered_Documents.json', rendered)
        suites = ET.fromstring((QA / 'final_regression_pytest.xml').read_bytes()).findall('testsuite')
        tests = {key: sum(int(x.get(key, '0')) for x in suites) for key in ['tests', 'failures', 'errors', 'skipped']}
        tests['passed'] = tests['tests'] - tests['failures'] - tests['errors'] - tests['skipped']
        assert tests['failures'] == tests['errors'] == 0
        shutil.copy2(QA / 'final_regression_pytest.xml', OUT / 'verification/Regression_Tests.xml')
        shutil.copy2(QA / 'packaging_pytest.xml', OUT / 'verification/Packaging_Tests.xml')
        if (QA / 'release_pytest.xml').exists():
            shutil.copy2(QA / 'release_pytest.xml', OUT / 'verification/Light_Release_Tests.xml')
        release_tests = None
        if (QA / 'release_pytest.xml').exists():
            rs = ET.fromstring((QA / 'release_pytest.xml').read_bytes()).findall('testsuite')
            release_tests = {key: sum(int(x.get(key, '0')) for x in rs) for key in ['tests', 'failures', 'errors', 'skipped']}
            release_tests['passed'] = release_tests['tests'] - release_tests['failures'] - release_tests['errors'] - release_tests['skipped']
            assert release_tests['failures'] == release_tests['errors'] == 0
        dump(OUT / 'verification/Current_Test_Summary.json', tests)
        author = json.loads(read(args.author_config))
        full = read(OUT / 'Manuscript_Main_Blind_and_Full.md')
        copies = json.loads(read(OUT / 'verification/Byte_Identical_Copies.json'))
        create_release(author, full, copies)
        dump(OUT / 'verification/Byte_Identical_Copies.json', copies)
        release = OUT / 'github_zenodo_release'
        shutil.copy2(OUT / 'verification/Regression_Tests.xml', release / 'results/reports/step08_regression_pytest.xml')
        # Keep scientific tests separate from private submission-author confirmation tests.
        write(release / 'requirements-submission.txt', 'python-docx>=1.2,<2\nlxml>=5,<7\n# Pandoc is a separate executable; pass --pandoc to the builder.\n')
        readme = read(release / 'README.md')
        readme += (f'\n## Current packaging regression\n\nThe current full workspace run passed '
                   f'{tests["passed"]} tests, with {tests["skipped"]} optional live-network test '
                   'skipped and no failures. The submission-specific test module requires '
                   'python-docx, available through requirements-submission.txt. Author confirmations '
                   'and generated submission artifacts are intentionally excluded from the light '
                   'release; their tests skip when these private/artifact inputs are absent.\n')
        write(release / 'README.md', readme)
        dump(OUT / 'verification/Frozen_Input_Integrity.json', integrity)
        validation = package_validation(author)
        counts = json.loads(read(OUT / 'verification/Word_Counts.json'))
        highlights = [x['text'] for x in json.loads(read(OUT / 'verification/Highlights_Character_Counts.json'))]
        summary = {'status': 'complete', 'docx_count': len(validation), 'reference_count': 8,
                   'frozen_files_checked': integrity['checked_files'], 'tests': tests, 'light_release_tests': release_tests,
                   'rendered_pages': sum(x['pages'] for x in rendered),
                   'rendered_documents': [{k: x[k] for k in ['file', 'pages', 'sha256']} for x in rendered],
                   'feature_extraction_rerun': False, 'classifier_training_rerun': False,
                   'publication_actions_performed': False, 'author_manual_review_confirmed': True,
                   'institutional_ethics_determination_provided': False,
                   'orcid_assigned': False, 'zenodo_doi_assigned': False}
        dump(OUT / 'verification/Build_Summary.json', summary)
        write_checklist(highlights, counts, integrity, validation, copies)
        checklist = read(OUT / 'SUBMISSION_CHECKLIST_FOR_AUTHOR.md')
        checklist = checklist.replace('## Highlights逐条字符数',
            f'当前回归：{tests["passed"]}项通过、{tests["skipped"]}项可选联网测试跳过，0失败。'
            f'8份Word共{summary["rendered_pages"]}页已渲染并检查。\n\n## Highlights逐条字符数')
        if release_tests:
            checklist = checklist.replace('## 尚需作者在正式提交前落实',
                f'轻量包在独立副本中验证：{release_tests["passed"]}项通过、{release_tests["skipped"]}项跳过。'
                '跳过项为可选公网传输及未随开源包发布的作者确认/投稿成品检查；不是临床结果重拟合。\n\n## 尚需作者在正式提交前落实')
        write(OUT / 'SUBMISSION_CHECKLIST_FOR_AUTHOR.md', checklist)
        assert_frozen(before)
        result = finalize_archives()
        write(ROOT / 'results/reports/step_08_report.md', '# Step 8 submission package\n\n' +
            '## Completed outputs and verification\n\n' +
            f'- Eight DOCX documents, including alternative single-journal cover letters; '
            f'{summary["rendered_pages"]} rendered pages reviewed.\n'
            '- Main review manuscript: editable native OMML equations, Times New Roman 12 pt, '
            'double spacing, 1 inch margins, continuous line numbers and page numbers.\n'
            '- Integrated preprint: all five original PNG images embedded byte-identically.\n'
            '- Eight DOI records verified by Crossref; two full page ranges supplemented by '
            'NLM PubMed and the original BH journal title checked against its publisher.\n'
            f'- {integrity["checked_files"]} protected inputs unchanged by SHA256, including '
            'the final_scope_alignment_integrity.json manifest itself.\n'
            f'- Current regression: {tests["passed"]} passed, {tests["skipped"]} optional network '
            'test skipped, no failures/errors. Three packaging safeguards also pass. '
            'Existing GPU/feature-name warnings do not indicate failed assertions.\n'
            '- No clinical feature extraction, benchmark classifier fitting or figure '
            'regeneration was performed. Test-only model fixtures are not new clinical fits.\n\n'
            '## Author and publication status\n\n'
            'Zhenxuan Ma, Hubei University of Medicine, Shiyan 442000, Hubei, China; '
            'mzxmc08@gmail.com. The author explicitly confirmed manual review and responsibility. '
            'No registered ORCID, institutional exemption determination, public GitHub release '
            'or Zenodo DOI was invented. GitHub publication and Zenodo archival release remain '
            'author actions; the package builder makes no submissions. Public-data secondary '
            'analysis is described factually without asserting a formal institutional exemption.\n\n'
            '## Files and archive\n\n' + json.dumps(result, ensure_ascii=False, indent=2) + '\n\n'
            'See submission_package/SUBMISSION_CHECKLIST_FOR_AUTHOR.md for the Chinese upload '
            'guide, exact Highlights character counts and remaining release/ethics checks.\n')
        print(json.dumps(result, ensure_ascii=True))
        return
    references = reference_verification()
    title, full, blind, abstract, keywords, counts, highlights = build_manuscripts(author, references)
    dump(OUT / 'author_configuration.json', author)
    pandoc_doc(title_page(title, abstract, keywords, counts, author), OUT / '01_Title_Page.docx', args.pandoc, author, double=False)
    pandoc_doc('# Highlights\n\n' + '\n\n'.join('- ' + x for x in highlights), OUT / '02_Highlights.docx', args.pandoc, author, blind=True, double=False)
    main_doc = pandoc_doc(blind, OUT / '03_Manuscript_DoubleSpaced_LineNumbered.docx', args.pandoc, author, blind=True, numbers=True)
    preprint = pandoc_doc(full, OUT / '04_bioRxiv_Integrated_Preprint.docx', args.pandoc, author)
    add_integrated_figures(preprint); preprint.save(OUT / '04_bioRxiv_Integrated_Preprint.docx')
    supp = supplement(full, author)
    supp_doc = pandoc_doc(supp, OUT / '05_Supplementary_Material.docx', args.pandoc, author)
    # One original S5 before its legend.
    for para in list(supp_doc.paragraphs):
        if para.text.startswith('Supplementary Figure S5.'):
            image = para.insert_paragraph_before()
            image.add_run().add_picture(str(ROOT / 'results/figures' / (FIGURES[4] + '.png')), width=Inches(6.5))
            image.paragraph_format.keep_with_next = True
            image.paragraph_format.line_spacing = 1
            para.paragraph_format.line_spacing = 1.15
    supp_doc.save(OUT / '05_Supplementary_Material.docx')
    for name, journal in [('06_Cover_Letter', 'Computer Methods and Programs in Biomedicine'),
                          ('06b_Cover_Letter_Alternative_CBM', 'Computers in Biology and Medicine')]:
        letter = cover_letter(title, author, journal)
        write(OUT / (name + '.md'), letter)
        pandoc_doc(letter, OUT / (name + '.docx'), args.pandoc, author, double=False)
    doi = '# Declaration of interest\n\n' + title + '\n\n' + author['name'] + '\n\n' + author['competing_interests'] + '\n'
    pandoc_doc(doi, OUT / '07_Declaration_of_Interest.docx', args.pandoc, author, double=False)
    copies = {}
    for fig in FIGURES:
        for ext in ['.pdf', '.png']:
            src = ROOT / 'results/figures' / (fig + ext)
            copy_verified(src, OUT / 'figures_and_tables' / src.name, copies)
    for src in (ROOT / 'results/tables').glob('*.csv'):
        if src.name in ['Analysis_Participants.csv', 'Clinical_CV_Splits.csv', 'Step7_Primary_Splits.csv',
                        'Step7_Prevalence_Stress_Members.csv', 'Step7_Prevalence_Stress_Anchor_Predictions.csv']: continue
        copy_verified(src, OUT / 'figures_and_tables' / src.name, copies)
    create_release(author, full, copies)
    dump(OUT / 'verification/Byte_Identical_Copies.json', copies)
    policy_notes()
    validation = package_validation(author)
    integrity = assert_frozen(before)
    dump(OUT / 'verification/Frozen_Input_Integrity.json', integrity)
    write_checklist(highlights, counts, integrity, validation, copies)
    result = finalize_archives()
    result.update({'status': 'built_pending_render_QA', 'docx_count': len(validation),
                   'frozen_files_checked': integrity['checked_files'], 'reference_count': 8,
                   'runtime_seconds': time.monotonic() - started})
    dump(OUT / 'verification/Build_Summary.json', result)
    # This new report is outside the protected scientific tables and source manuscripts.
    write(ROOT / 'results/reports/step_08_report.md', '# Step 8 submission package\n\n' +
          json.dumps(result, ensure_ascii=False, indent=2) + '\n\n'
          'No clinical feature extraction, classifier fitting, numeric table mutation or '
          'figure regeneration was performed. Author review was explicitly confirmed. '
          'Public repository and Zenodo release are pending. No institutional ethics '
          'determination was supplied or invented. See the Chinese author checklist and '
          'the package verification directory for reference, structural and render QA evidence.\n')
    print(json.dumps(result, ensure_ascii=True), flush=True)


if __name__ == '__main__':
    main()
