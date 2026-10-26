"""Finalize submission typography and verified bibliography without clinical refitting.

--figure-only uses the analysis environment. --build/--finalize use the bundled
document runtime. Historical scientific outputs and the v0.1.0 archive stay fixed.
"""
from __future__ import annotations
import argparse,copy,csv,hashlib,importlib.util,json,re,shutil,subprocess,sys,zipfile
from pathlib import Path

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission_package';WORK=ROOT/'work/step08_1'
ARCHIVE=ROOT/'results/archive_before_step08_1'
spec=importlib.util.spec_from_file_location('pack08',ROOT/'scripts/08_build_submission_package.py')
pack=importlib.util.module_from_spec(spec);spec.loader.exec_module(pack)
AFFILIATION='School of Clinical Medicine, Hubei University of Medicine, Shiyan 442000, Hubei, China'
ORCID='0009-0004-5432-4999'
FIGURE_NAME='Figure3_LSN_Debiasing_Performance_Typeset'

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def dump(p,obj):pack.dump(p,obj)
def read(p):return Path(p).read_text(encoding='utf-8')
def frozen_check():
    original=json.loads(read(ROOT/'work/step09/scientific_baseline.json'))
    for rel,digest in original.items():assert sha(ROOT/rel)==digest,rel
    return original
def orcid_valid(value):
    d=value.replace('-','');total=0
    for c in d[:-1]:total=(total+int(c))*2
    check=(12-total%11)%11
    return len(d)==16 and d[-1]==('X' if check==10 else str(check))
def archive():
    if (ARCHIVE/'archive_manifest.json').exists():return
    files=list(OUT.glob('*.docx'))+list(OUT.glob('*.md'))+[OUT/'04_bioRxiv_Integrated_Preprint.pdf']
    files+=list((OUT/'verification').glob('*.*'))
    files+=[ROOT/'results/MANUSCRIPT_FULL_TEXT.md',ROOT/'results/reports/MANUSCRIPT_MASTER_DRAFT.md',ROOT/'config/submission_author.json']
    hashes={}
    for p in files:
        target=ARCHIVE/p.relative_to(ROOT);target.parent.mkdir(parents=True,exist_ok=True)
        shutil.copy2(p,target);hashes[str(target.relative_to(ARCHIVE))]=sha(target)
    dump(ARCHIVE/'archive_manifest.json',hashes)

def figure_only():
    # Source code is evaluated unchanged except removal of one tiny text artist.
    # All numerical artists are compared between the original and typeset figure.
    import numpy as np,pandas as pd,matplotlib.pyplot as plt
    sys.path.insert(0,str(ROOT))
    from src.common import load_config
    source=read(ROOT/'src/multiscale_plotting.py')
    line="    roc.text(0,-.22,annotation,transform=roc.transAxes,fontsize=6)"
    assert source.count(line)==1
    cfg=load_config(str(ROOT/'config/config.yaml'));tables=ROOT/'results/tables'
    args=[cfg,pd.read_parquet(tables/'Step7_OOF_Predictions.parquet'),
          pd.read_csv(tables/'Table3_Benchmark_Metrics_Full.csv'),pd.read_csv(tables/'Step7_Paired_Significance.csv'),
          pd.read_csv(tables/'Analysis_Cohort_Counts.csv'),pd.read_csv(tables/'Step7_AUC_Pair_Decomposition.csv'),
          pd.read_csv(tables/'Step7_Prevalence_Stress_Summary.csv')]
    signatures=[]
    def capture(fig,folder,name,dpi):
        if name=='Figure3_LSN_Debiasing_Performance':
            payload=[]
            for ax in fig.axes:
                payload.extend([np.asarray(line.get_xydata(),dtype=float).tobytes().hex() for line in ax.lines])
                payload.extend([np.asarray(p.get_path().vertices,dtype=float).tobytes().hex() for p in ax.patches])
            signatures.append(hashlib.sha256(json.dumps(payload).encode()).hexdigest())
            if len(signatures)==2:
                folder=OUT/'figures_and_tables';folder.mkdir(exist_ok=True)
                fig.savefig(folder/(FIGURE_NAME+'.pdf'),dpi=300,bbox_inches='tight')
                fig.savefig(folder/(FIGURE_NAME+'.png'),dpi=300,bbox_inches='tight')
        plt.close(fig)
    for code in [source,source.replace(line,'    # DeLong comparisons are set in the full-size figure legend.')]:
        ns={'__name__':'src.step08_1_figure','__package__':'src'}
        exec(compile(code,'src/multiscale_plotting.py','exec'),ns)
        ns['save']=capture;ns['step7_figures'](*args)
    assert signatures[0]==signatures[1]
    dump(WORK/'figure3_numeric_artist_verification.json',{'status':'passed',
         'original_and_typeset_numeric_artist_sha256':signatures,
         'change':'Removed one sub-axis text annotation; exact p and q values moved to legend',
         'frozen_original_pdf_sha256':sha(ROOT/'results/figures/Figure3_LSN_Debiasing_Performance.pdf'),
         'typeset_pdf_sha256':sha(OUT/'figures_and_tables'/(FIGURE_NAME+'.pdf'))})
    frozen_check();print('Figure 3 numerical geometry unchanged; original frozen image preserved.')

def bibliography():
    refs=json.loads(read(WORK/'references_verified.json'))
    assert len(refs)==28 and len({r['doi'].lower() for r in refs})==28
    for row in refs:
        raw=json.loads(read(WORK/'crossref'/f'{row["reference"]:02}.json'))['message']
        assert raw['DOI'].lower()==row['doi'].lower()
        if row['reference']==20:
            row['pages_or_article']='1-22'
            row['supplement_source']='https://www.jstatsoft.org/citationstylelanguage/get/associacao-brasileira-de-normas-tecnicas?publicationId=1756&submissionId=1757'
        if row['reference']==22:row['title']+=': '+pack.clean(raw['subtitle'][0])
        if row['reference']==25:
            row['journal']=raw['container-title'][-1]
            row['publisher']=raw['publisher'];row['publisher_location']=raw['publisher-location']
            row['editor']='Csurka G'
        if row['reference']==27:
            authors=[a for a in raw['author'] if 'family' in a]
            row['authors']=', '.join(a['family']+' '+pack.initials(a.get('given','')) for a in authors)
            row['authors']+='; on behalf of the STRATOS Topic Group on diagnostic tests and prediction models'
        if row['reference'] in [2,7]:row['supplement_source']='NLM PubMed pagination verified in original package'
        row['crossref_sha256']=sha(WORK/'crossref'/f'{row["reference"]:02}.json')
        row['accessed']='2026-10-26'
        assert row['pages_or_article'] and row['journal']
        if row['type']=='journal-article':assert row['volume'] and row['issue']
    return refs

def revise_text(original):
    full=original.replace('**Affiliation:** Hubei University of Medicine','**Affiliation:** School of Clinical Medicine, Hubei University of Medicine')
    full=re.sub(r'\*\*ORCID:\*\*[^\n]*','**ORCID:** https://orcid.org/'+ORCID,full)
    full=re.sub(r'## Highlights\n.*?(?=## References\n)','',full,flags=re.S)
    # Relevant scholarship contextualizes the existing findings; no new measured result is introduced.
    old='alongside sequence-variant assays [1–3].'
    full=full.replace(old,'alongside sequence-variant assays [1–3,9]. Fragment-size selection, transcription-associated coverage and lung-cancer fragmentome studies establish complementary uses of these signals [10,11,13]. Griffin further links nucleosome protection profiles with cancer phenotypes while explicitly addressing GC bias [14].')
    full=full.replace('Study Source combines specimen processing, sequencing procedures, recruitment and population composition.',
        'Study Source combines specimen processing, sequencing procedures, recruitment and population composition. This joint structure is consistent with the broader challenge of batch effects in high-throughput molecular data [15,16].')
    full=full.replace('Window counts, IQR and the global integer-length frequency distribution were retained in compact local tables.',
        'Window counts, IQR and the global integer-length frequency distribution were retained in compact local tables. Descriptive UMAP embeddings used the published manifold-learning implementation [23]; NumPy supported array operations [28].')
    full=full.replace('999 restricted Freedman–Lane residual permutations.','999 restricted Freedman–Lane residual permutations [18,19].')
    full=full.replace('learning regularized coefficients for the two correlated blocks.',
        'learning regularized coefficients for the two correlated blocks. Elastic-net regularization and the generalized-linear-model regularization path provide the methodological context for this penalized comparison [21,20].')
    full=full.replace('This analysis evaluated source-only scale selection separately from fixed Hybrid.',
        'This analysis evaluated source-only scale selection separately from fixed Hybrid. All preprocessing and selection remained inside the source split, addressing the training/test separation failures documented in leakage audits [17].')
    full=full.replace('DANN used training-source domains exclusively.',
        'Random Forest and XGBoost follow their established ensemble formulations [24,22]. The gradient-reversal mechanism follows domain-adversarial representation learning [25]; here DANN used training-source domains exclusively.')
    full=full.replace('AUPRC is interpreted together with case proportion because its baseline changes with prevalence.',
        'AUPRC is interpreted together with case proportion because its baseline changes with prevalence [26]. Calibration was evaluated separately from ranking, following the distinction emphasized in prediction-model evaluation [27].')
    full=full.replace('The observed target-specific comparisons identify the conditions in which the fused representation changes ranking and calibration;',
        'Published low-coverage copy-number methods such as ichorCNA provide a separate route to estimating tumor fraction and broad genomic imbalance [12]. Fragmentation features therefore warrant comparison with matched copy-number measurements rather than automatic interpretation as CNA. The observed target-specific comparisons identify the conditions in which the fused representation changes ranking and calibration;')
    full=full.replace('other covariate-aware or reference-batch ComBat designs remain separate methods [4].',
        'other covariate-aware or reference-batch ComBat designs remain separate methods [4,15,16].')
    full=full.replace('patient-wise self-normalization as a direct way to remove companion-dependent target preprocessing.',
        'patient-wise self-normalization as a direct way to remove companion-dependent target preprocessing.')
    full=full.replace('**Figure 3. Multiscale representation and fixed-patient composition dependence.**',
        '**Figure 3. Multiscale representation and fixed-patient composition dependence.**')
    sig=pack.rows('Step7_Paired_Significance.csv')
    compare=[r for r in sig if r['cohort_scope']=='Primary_Bilateral' and r['model']=='elasticnet' and r['method']=='Hybrid_MultiScale' and r['comparator']=='Global_Zscore']
    values='; '.join(f'{r["center"].split(" et al")[0]}: p={float(r["p"]):.6f}, BH q={float(r["q"]):.6f}' for r in compare)
    full=full.replace('(c) Primary calibration curves and Brier scores.',
        f'Target-specific Hybrid-versus-Global DeLong tests were {values}. These comparisons concern the separate held-out studies; the pooled curves are descriptive. (c) Primary calibration curves and Brier scores.')
    # Every original table string, including all numbers and row order, is invariant.
    assert [x for x in full.splitlines() if x.startswith('|')]==[x for x in original.splitlines() if x.startswith('|')]
    assert not re.search(r'\bAl\b',full)
    return full

def renumber_and_format(full,refs):
    body=full.split('## References\n')[0];order=[]
    pattern=re.compile(r'\[(\d+(?:\s*[–-]\s*\d+)?(?:\s*,\s*\d+(?:\s*[–-]\s*\d+)?)*)\]')
    def numbers(s):
        out=[]
        for token in s.split(','):
            bounds=re.split('[–-]',token)
            out.extend(range(int(bounds[0]),int(bounds[-1])+1))
        return out
    for match in pattern.finditer(body):
        for n in numbers(match.group(1)):
            if n not in order:order.append(n)
    assert set(order)==set(range(1,29)),set(range(1,29))-set(order)
    mapping={old:new for new,old in enumerate(order,1)}
    def replace(m):return '['+','.join(str(n) for n in sorted(mapping[x] for x in numbers(m.group(1))))+']'
    body=pattern.sub(replace,body);lookup={r['reference']:r for r in refs};verified=[];formatted=[]
    for old in order:
        r=copy.deepcopy(lookup[old]);r['source_reference_id']=old;r['reference']=mapping[old]
        first=f'{r["reference"]}. {r["authors"]}. {r["title"]}. '
        if r['type']=='book-chapter':
            end=f'In: {r["editor"]}, editor. {r["journal"]}. {r["publisher_location"]}: {r["publisher"]}; {r["year"]}. p. {r["pages_or_article"]}.'
        elif r['type']=='proceedings-article':end=f'{r["journal"]}. {r["year"]}. p. {r["pages_or_article"]}.'
        else:end=f'{r["journal"]}. {r["year"]};{r["volume"]}'+(f'({r["issue"]})' if r['issue'] else '')+f':{r["pages_or_article"]}.'
        formatted.append(first+end+' https://doi.org/'+r['doi']);verified.append(r)
    return body+'## References\n\n'+'\n\n'.join(formatted)+'\n',verified

def complete_tables(text):
    """One complete table per source block, with all numeric strings unchanged."""
    lines=text.splitlines();out=[];i=0;number=0
    while i<len(lines):
        if not lines[i].lstrip().startswith('|'):out.append(lines[i]);i+=1;continue
        block=[]
        while i<len(lines) and lines[i].lstrip().startswith('|'):block.append(lines[i]);i+=1
        data=[[v.strip() for v in x.strip().strip('|').split('|')] for x in block]
        if len(data)<3:out+=block;continue
        number+=1;headers,rows=data[0],data[2:]
        name,caption=pack.TABLE_CAPTIONS.get(number,(f'Table {number}','Corresponding frozen results'))
        if headers==['pipeline','auroc','auroc_low','auroc_high']:name,caption='Supplementary Table S4','Secondary Combined LOCO baseline comparison'
        out+=['',f'**{name}. {caption}.**','']
        labels=[pack.HEADER_ALIASES.get(h,h.replace('_',' ')) for h in headers]
        out+=['| '+' | '.join(labels)+' |','| '+' | '.join(['---']*len(headers))+' |']
        for row in rows:
            cells=[]
            for value in row:
                v=pack.METHOD_ALIASES.get(value,value)
                if v.startswith(('Cristiano et al','Jiang et al','Snyder et al')):v=v.split(' et al')[0]
                v=v.replace('Primary_Bilateral','Primary bilateral').replace('External_HighBurden','External high burden').replace('Combined_LOCO','Combined LOCO')
                cells.append(v)
            out.append('| '+' | '.join(cells)+' |')
        out.append('')
    return '\n'.join(out)+'\n'

def table_layout(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    def section_end(where,landscape=False,before=False):
        p=OxmlElement('w:p');ppr=OxmlElement('w:pPr');sect=copy.deepcopy(doc.sections[-1]._sectPr)
        size=sect.find(qn('w:pgSz'));size.set(qn('w:w'),'15840' if landscape else '12240');size.set(qn('w:h'),'12240' if landscape else '15840')
        size.set(qn('w:orient'),'landscape' if landscape else 'portrait')
        typ=sect.find(qn('w:type'))
        if typ is None:typ=OxmlElement('w:type');sect.insert(0,typ)
        typ.set(qn('w:val'),'nextPage');ppr.append(sect);p.append(ppr)
        if before:where.addprevious(p)
        else:where.addnext(p)
    for table in doc.tables:
        n=len(table.columns);wide=n>=9;width=648 if wide else 468
        if n==11:weights=[74,54,47,47,47,47,52,52,80,43,43]
        elif n==9:weights=[90,55,65,65,55,75,75,50,50]
        elif n==7:weights=[115,36,54,72,58,63,70]
        elif n==6:weights=[110,65,75,75,65,85]
        else:weights=[130]+[80]*(n-1)
        widths=[width*x/sum(weights) for x in weights]
        for c,w in zip(table.columns,widths):c.width=Pt(w)
        # One authoritative table-level margin/layout definition.
        for x in list(table._tbl.tblPr.findall(qn('w:tblCellMar'))):table._tbl.tblPr.remove(x)
        margins=OxmlElement('w:tblCellMar')
        for side,val in [('top',32),('bottom',32),('left',35),('right',35)]:
            el=OxmlElement('w:'+side);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
        table._tbl.tblPr.append(margins)
        for ri,row in enumerate(table.rows):
            tr=row._tr.get_or_add_trPr()
            if tr.find(qn('w:cantSplit')) is None:tr.append(OxmlElement('w:cantSplit'))
            if ri==0 and tr.find(qn('w:tblHeader')) is None:tr.append(OxmlElement('w:tblHeader'))
            for ci,cell in enumerate(row.cells):
                cell.width=Pt(widths[ci])
                for p in cell.paragraphs:
                    if ri==0:
                        mapping={'AUROC lower':'AUROC\nlower','AUROC upper':'AUROC\nupper',
                            'Specificity shift (pp)':'Specificity\nshift (pp)','Sensitivity lower':'Sensitivity\nlower',
                            'Sensitivity upper':'Sensitivity\nupper','Locked threshold':'Locked\nthreshold','ICC bypassed':'ICC\nbypassed'}
                        if p.text in mapping:p.text=mapping[p.text]
                    p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(1)
                    p.paragraph_format.space_before=Pt(0)
                    # Each compact table fits on a page: keep its rows together.
                    p.paragraph_format.keep_with_next=ri<len(table.rows)-1
                    for run in p.runs:pack.set_font(run,8.5 if n>=7 else 9);run.bold=ri==0
        if wide:
            caption=table._tbl.getprevious();assert caption.tag==qn('w:p')
            start=caption
            end=table._tbl
            caption_text=''.join(caption.itertext())
            if 'Table 3A.' in caption_text:
                # Place its short section introduction on the same landscape page.
                candidate=caption.getprevious()
                while candidate is not None and candidate.tag==qn('w:p'):
                    visible=''.join(candidate.itertext())
                    if '3.2. Primary bilateral generalization' in visible:start=candidate;break
                    candidate=candidate.getprevious()
            if 'Table 3C.' in caption_text:
                previous=caption.getprevious();following=table._tbl.getnext()
                if previous is not None and 'The observed healthy outcome' in ''.join(previous.itertext()):start=previous
                if following is not None and 'For fixed ElasticNet' in ''.join(following.itertext()):end=following
            section_end(start,before=True);section_end(end,landscape=True)
    # A copied start=1 restarts LibreOffice numbering despite restart=continuous.
    for index,sect in enumerate(doc._element.body.iter(qn('w:sectPr'))):
        line=sect.find(qn('w:lnNumType'))
        if line is not None and index>0:line.attrib.pop(qn('w:start'),None)
    return doc

def add_figures(doc):
    from docx.shared import Inches
    for p in list(doc.paragraphs):
        if p.text=='Figure legends':p._element.getparent().remove(p._element)
    for p in list(doc.paragraphs):
        match=re.match(r'^(Figure ([1-4])\.|Supplementary Figure S5\.)',p.text)
        if not match:continue
        index=int(match.group(2))-1 if match.group(2) else 4
        path=OUT/'figures_and_tables'/(FIGURE_NAME+'.png') if index==2 else ROOT/'results/figures'/(pack.FIGURES[index]+'.png')
        image=p.insert_paragraph_before();image.paragraph_format.page_break_before=True
        image.paragraph_format.keep_with_next=True;image.paragraph_format.line_spacing=1
        image.add_run().add_picture(str(path),width=Inches(6.5))
        p.paragraph_format.line_spacing=1.15
        for run in p.runs:pack.set_font(run,11)

def build(args):
    before=frozen_check();archive();assert orcid_valid(ORCID)
    author=json.loads(read(ROOT/'config/submission_author.json'));author.update(affiliation=AFFILIATION,orcid=ORCID)
    receipt=json.loads(read(WORK/'zenodo_live_record.json'));assert receipt['doi']==author['zenodo_doi']
    full,refs=renumber_and_format(revise_text(read(ARCHIVE/'results/MANUSCRIPT_FULL_TEXT.md')),bibliography())
    blind=re.sub(r'\*\*Author:\*\*.*?(?=## Abstract)','',full,flags=re.S).replace(author['name'],'The sole author')
    begin=blind.index('Code and compact statistical outputs are distributed at ');end=blind.index('The release contains',begin)
    blind=blind[:begin]+'The versioned code and statistical archive are supplied through the separate title page for anonymous review. '+blind[end:]
    assert not any(x in blind for x in [ORCID,author['email'],author['name'],'mzxuan08',AFFILIATION])
    title=full.splitlines()[0][2:];abstract=pack.extract_section(full,'## Abstract','## 1. Introduction')
    abstract_text,keywords=abstract.split('**Keywords:**');body=full.split('## 1. Introduction\n')[1].split('## Data and code availability\n')[0]
    counts={'abstract_words':pack.word_count(abstract_text),'main_body_words_excluding_tables':pack.word_count('\n'.join(x for x in body.splitlines() if not x.startswith('|'))),
            'main_body_including_tables_words':pack.word_count(body),'full_markdown_words':pack.word_count(full),
            'definition':'Whitespace-delimited tokens containing a letter or numeral; body is Introduction through Conclusions, excludes declarations, references, legends and Highlights.'}
    dump(ROOT/'config/submission_author.json',author);dump(OUT/'author_configuration.json',author)
    for p in [OUT/'Manuscript_Main_Blind_and_Full.md',ROOT/'results/MANUSCRIPT_FULL_TEXT.md',ROOT/'results/reports/MANUSCRIPT_MASTER_DRAFT.md']:pack.write(p,full)
    pack.write(OUT/'Manuscript_Main_Blinded.md',blind);dump(OUT/'verification/Word_Counts.json',counts)
    dump(OUT/'verification/reference_verification.json',{'status':'passed','references':refs,'Crossref_verified_DOIs':28,
        'pagination_supplements':['Jiang and DeLong: NLM PubMed','Friedman: publisher JSS 1-22'],
        'note':'Volume/issue fields are not invented for book chapters or conference proceedings; DANN is the DOI-bearing 2017 Springer chapter, not mislabeled as the 2016 JMLR article.'})
    with (OUT/'verification/References_Verified.csv').open('w',encoding='utf-8',newline='') as f:
        fields=sorted(set().union(*(r.keys() for r in refs)));w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(refs)
    shutil.copytree(WORK/'crossref',OUT/'verification/crossref_step08_1',dirs_exist_ok=True)
    pack.write(OUT/'verification/Reference_Verification_Notes.md',
        '# Verified bibliography\n\n28 distinct DOIs were matched to Crossref works records and numbered by first citation. 20 references were added. Missing journal pagination was supplemented from NLM PubMed or the original JSS publisher citation. Conference and book records have no invented journal volumes/issues. The DANN DOI refers to the 2017 Springer chapter. The original scikit-learn and PyTorch papers are not assigned fabricated Crossref DOIs; the bibliography uses DOI-bearing primary methods/software references.\n')
    pack.presentation_tables=complete_tables
    original_title=pack.title_page(title,abstract_text.strip(),keywords.strip(),counts,author)
    original_title=original_title.replace('ORCID: Not registered.','ORCID: https://orcid.org/'+ORCID)
    original_title=re.sub(r'## Data and code release\n.*','## Data and code release\n\nPublic repository: '+author['repository_url']+'. Immutable v0.1.0 software archive: https://doi.org/'+author['zenodo_doi']+'.\n',original_title,flags=re.S)
    def create(text,name,**kw):
        doc=pack.pandoc_doc(text,OUT/name,args.pandoc,author,**kw)
        if name.startswith('03_'):
            from docx.shared import Pt
            for paragraph in doc.paragraphs:
                paragraph.paragraph_format.space_before=Pt(0)
                paragraph.paragraph_format.space_after=Pt(4)
                if paragraph.style.name.startswith('Heading'):
                    paragraph.paragraph_format.space_before=Pt(6)
        table_layout(doc);doc.save(OUT/name);return doc
    create(original_title,'01_Title_Page.docx',double=False)
    highlights=json.loads(read(OUT/'verification/Highlights_Character_Counts.json'))
    assert all(len(h['text'])<=85 for h in highlights)
    create('# Highlights\n\n'+'\n\n'.join('- '+h['text'] for h in highlights),'02_Highlights.docx',blind=True,double=False)
    create(blind,'03_Manuscript_DoubleSpaced_LineNumbered.docx',blind=True,numbers=True)
    doc=create(full,'04_bioRxiv_Integrated_Preprint.docx');add_figures(doc);doc.save(OUT/'04_bioRxiv_Integrated_Preprint.docx')
    supp=pack.supplement(full.replace('## References','## Highlights\n\n## References'),author)
    supp=supp.replace('This supplement accompanies the frozen analysis and documents the scope of the statistical tables. It introduces no new fitted predictions or observations.',
        'This supplement decomposes score comparability and documents the statistical tables supporting the normalization-scale and fixed-anchor analyses.')
    doc=create(supp,'05_Supplementary_Material.docx')
    from docx.shared import Inches,Pt
    for p in list(doc.paragraphs):
        if p.text.startswith('Supplementary Figure S5.'):
            image=p.insert_paragraph_before();image.add_run().add_picture(str(ROOT/'results/figures'/(pack.FIGURES[4]+'.png')),width=Inches(6.5))
            image.paragraph_format.keep_with_next=True;image.paragraph_format.line_spacing=1;p.paragraph_format.line_spacing=1.15
    doc.save(OUT/'05_Supplementary_Material.docx')
    for name in ['06_Cover_Letter','06b_Cover_Letter_Alternative_CBM']:
        text=read(ARCHIVE/'submission_package'/(name+'.md')).replace('Hubei University of Medicine, Shiyan 442000, Hubei, China',AFFILIATION)
        text=re.sub(re.escape(author['email'])+r'\s*$',author['email']+' | ORCID: '+ORCID+'\n',text)
        pack.write(OUT/(name+'.md'),text)
        doc=create(text,name+'.docx',double=False)
        for p in doc.paragraphs:p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(4)
        doc.save(OUT/(name+'.docx'))
    create('# Declaration of interest\n\n'+title+'\n\n'+author['name']+'\n\n'+author['competing_interests'],'07_Declaration_of_Interest.docx',double=False)
    dump(WORK/'build_preflight.json',{'status':'passed','counts':counts,'references':28,'author':author,'protected_files':len(before)})
    validate_package();assert frozen_check()==before
    print(json.dumps({'status':'built','references':len(refs),'counts':counts},ensure_ascii=True))

def validate_package():
    from docx import Document
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    from lxml import etree
    from docx.oxml.ns import qn
    refs=json.loads(read(OUT/'verification/reference_verification.json'))['references']
    assert len(refs)==28
    result={}
    for p in sorted(OUT.glob('*.docx')):
        doc=Document(p);text='\n'.join(x.text for x in doc.paragraphs)
        assert not re.search(r'\bAl\b|Not registered|PENDING_RELEASE',text)
        with zipfile.ZipFile(p) as z:
            assert z.testzip() is None;xml=etree.fromstring(z.read('word/document.xml'))
            if p.name.startswith(('03_','04_')):
                assert all(x.text!='Highlights' for x in doc.paragraphs)
                assert len(doc.tables)==7,len(doc.tables)
                assert [len(t.columns) for t in doc.tables]==[5,5,2,11,7,6,9]
                assert not re.search(r'Table (?:3[ABC]|S3).*panel [abc]',text)
                for table in doc.tables:
                    assert table.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
                    for row in table.rows:assert row._tr.trPr.find(qn('w:cantSplit')) is not None
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:assert 8.5<=run.font.size.pt<=9
                assert len(xml.findall('.//m:oMath',ns))>=8
            if p.name.startswith('03_'):
                assert ORCID not in text and 'mzxuan08' not in text and 'Zhenxuan Ma' not in text
                assert xml.find('.//w:lnNumType',ns).get(qn('w:restart'))=='continuous'
            if p.name.startswith('04_'):
                assert ORCID in text and AFFILIATION in text and '10.5281/zenodo.23653766' in text
                actual={hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/')}
                expected={sha(ROOT/'results/figures'/(x+'.png')) for i,x in enumerate(pack.FIGURES) if i!=2}
                expected.add(sha(OUT/'figures_and_tables'/(FIGURE_NAME+'.png')));assert actual==expected
            result[p.name]={'tables':len(doc.tables),'native_equations':len(xml.findall('.//m:oMath',ns)),'sha256':sha(p),'status':'passed'}
    assert len(result)==8;dump(OUT/'verification/Document_Structural_Validation.json',result);return result

def finalize():
    before=frozen_check();validate_package()
    from pypdf import PdfReader
    status=json.loads(read(ROOT/'work/step08/render_status.json'))
    for x in status:assert x['returncode']==0 and x['sha256']==sha(OUT/x['file'])
    pdfsource=ROOT/'work/step08/render/04_bioRxiv_Integrated_Preprint/04_bioRxiv_Integrated_Preprint.pdf'
    shutil.copy2(pdfsource,OUT/'04_bioRxiv_Integrated_Preprint.pdf')
    reader=PdfReader(OUT/'04_bioRxiv_Integrated_Preprint.pdf');text='\n'.join(p.extract_text() or '' for p in reader.pages)
    assert 'Highlights' not in text and ORCID in text and '10.5281/zenodo.23653766' in ''.join(text.split())
    assert not re.search(r'\bAl\b|Not registered|panel [abc]\. Primary bilateral',text)
    compact=''.join(text.split())
    for row in json.loads(read(OUT/'verification/reference_verification.json'))['references']:
        assert row['doi'] in compact,row['doi']
    release=OUT/'github_zenodo_release'
    for source,target in [(ROOT/'scripts/08_1_final_typesetting.py',release/'scripts/08_1_final_typesetting.py'),
         (ROOT/'results/MANUSCRIPT_FULL_TEXT.md',release/'results/MANUSCRIPT_FULL_TEXT.md'),
         (ROOT/'tests/test_submission_package.py',release/'tests/test_submission_package.py')]:shutil.copy2(source,target)
    released_readme=read(release/'README.md').replace('Zhenxuan Ma · Hubei University of Medicine, Shiyan 442000, Hubei, China','Zhenxuan Ma · '+AFFILIATION+'\n\nORCID: https://orcid.org/'+ORCID)
    released_readme=released_readme.replace('The recorded earlier research run had\n73 passing tests and one optional live-network test skipped; current package\nvalidation is reported separately and must not be described as a new full\nclinical extraction or benchmark run.',
        'The Step 8.1 verification run had 80 passing tests and one optional\nlive-network test skipped. This verified software and publication formatting;\nit did not repeat clinical extraction or benchmark fitting.')
    pack.write(release/'README.md',released_readme)
    cff=json.loads(read(release/'CITATION.cff'));cff['authors'][0].update(affiliation=AFFILIATION,orcid='https://orcid.org/'+ORCID);dump(release/'CITATION.cff',cff)
    zen=json.loads(read(release/'.zenodo.json'));zen['creators'][0].update(affiliation=AFFILIATION,orcid=ORCID);dump(release/'.zenodo.json',zen)
    example=json.loads(read(release/'config/submission_author.example.json'))
    example.update(affiliation=AFFILIATION,orcid=ORCID,zenodo_doi='10.5281/zenodo.23653766')
    dump(release/'config/submission_author.example.json',example)
    shutil.copy2(OUT/'verification/References_Verified.csv',release/'results/reports/References_Verified.csv')
    preflight=json.loads(read(WORK/'build_preflight.json'))
    summary={'status':'complete','step':'8.1','references':28,'references_added':20,'pdf_pages':len(reader.pages),
        'rendered_pages':sum(x['pages'] for x in status),'protected_scientific_files':len(before),'scientific_files_changed':0,
        'orcid':ORCID,'orcid_source':'supplied by author, ISO checksum verified','affiliation':AFFILIATION,
        'zenodo_doi':'10.5281/zenodo.23653766','zenodo_live_public_record_verified':True,
        'table_columns':[5,5,2,11,7,6,9],'wide_tables':'complete landscape tables, 8.5 pt, 1-inch margins',
        'figure3':'new typeset copy; original frozen figure retained; numerical artists identical',
        'counts':preflight['counts'],'pdf_sha256':sha(OUT/'04_bioRxiv_Integrated_Preprint.pdf')}
    import xml.etree.ElementTree as ET
    suites=list(ET.parse(WORK/'pytest.xml').getroot().iter('testsuite'))
    tests={key:sum(int(s.get(key,0)) for s in suites) for key in ['tests','failures','errors','skipped']}
    tests['passed']=tests['tests']-tests['skipped'];assert tests['failures']==tests['errors']==0
    summary['tests']=tests
    dump(OUT/'verification/Current_Test_Results.json',tests)
    dump(OUT/'verification/Current_Test_Summary.json',dict(tests,scope='Current full project tests after Step 8.1, no clinical extraction or prediction refitting'))
    shutil.copy2(WORK/'pytest.xml',OUT/'verification/Regression_Tests.xml')
    shutil.copy2(WORK/'packaging_tests.xml',OUT/'verification/Packaging_Tests.xml')
    dump(OUT/'verification/Build_Summary.json',summary)
    dump(OUT/'verification/Rendered_Documents.json',status)
    dump(OUT/'verification/Render_QA_Review.json',{'status':'passed','step':'8.1','rendered_documents':status,
        'method':'All final rendered pages inspected; complete tables, repeated headers, formulas, reference order and five integrated figures verified.'})
    dump(OUT/'verification/Editorial_Release_Integrity.json',{'status':'passed','protected_files':before,'changed_files':[],
        'figure3_typesetting_exception':'New derivative stored separately; frozen original remains byte-identical'})
    dump(OUT/'verification/Frozen_Input_Integrity.json',{'status':'passed','protected_files':before,'changed_files':[],'scope':'Scientific data and frozen original figures; authorized text and new figure derivative excluded'})
    shutil.copy2(WORK/'figure3_numeric_artist_verification.json',OUT/'verification/Figure3_Typesetting_Verification.json')
    copies=json.loads(read(OUT/'verification/Byte_Identical_Copies.json'))
    for rel,entry in list(copies.items()):
        if not (OUT/rel).exists() or sha(ROOT/entry['source'])!=sha(OUT/rel):del copies[rel]
        else:entry['sha256']=sha(ROOT/entry['source'])
    dump(OUT/'verification/Byte_Identical_Copies.json',copies)
    checklist=read(OUT/'SUBMISSION_CHECKLIST_FOR_AUTHOR.md')
    checklist=re.sub(r'作者：.*?。',f'作者：Zhenxuan Ma（马振轩）；{AFFILIATION}；mzxmc08@gmail.com；ORCID {ORCID}。',checklist,count=1)
    checklist=re.sub(r'04的PDF为\d+页',f'04的PDF为{len(reader.pages)}页',checklist)
    checklist=checklist.replace('- ORCID尚未注册；如平台强制要求请本人注册真实号码。','- ORCID已按作者提供的真实号码填写，并通过校验位检查。')
    checklist=checklist.replace('8条文献','28条文献')
    checklist=checklist.split('## 全部文件')[0]
    checklist+='## Step 8.1排版核验\n\n28篇文献全部匹配Crossref DOI；依首次出现顺序编号。Table 3A/3B/3C与S3均为完整三线表。宽表采用横向单独页面及8.5 pt表内字体，正文保持12 pt双倍行距。03/04不含Highlights。Figure 3b检验值移至完整图注，使用新排版副本；原始冻结图件仍保留。图件上传请选择带_Typeset后缀的Figure3文件。\n\n'
    inventory=sorted(str(p.relative_to(OUT)).replace('\\','/') for p in OUT.rglob('*') if p.is_file() and not any(x in p.parts for x in ['.git','__pycache__','.pytest_cache']))
    checklist+='## 全部文件\n\n'+'\n'.join('- `'+n+'`' for n in inventory)+'\n';pack.write(OUT/'SUBMISSION_CHECKLIST_FOR_AUTHOR.md',checklist)
    pack.write(ROOT/'results/reports/step_08_1_report.md','# Step 8.1 Final typesetting and reference expansion\n\n'+json.dumps(summary,indent=2,ensure_ascii=False)+'\n\nOriginal statistical CSV values, prediction files, feature matrices and all frozen figure hashes remain unchanged. The verified v0.1.0 Zenodo archive remains immutable. The new Figure 3 derivative removes only one small annotation and moves exact tests to the legend. Historical editorial versions are preserved under results/archive_before_step08_1/.\n')
    shutil.copy2(ROOT/'results/reports/step_08_1_report.md',release/'results/reports/step_08_1_report.md')
    assert frozen_check()==before
    pack.finalize_archives()
    print(json.dumps(summary,ensure_ascii=True))

def main():
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--figure-only',action='store_true');p.add_argument('--build',action='store_true');p.add_argument('--finalize',action='store_true');p.add_argument('--pandoc',type=Path)
    args=p.parse_args();WORK.mkdir(exist_ok=True,parents=True)
    if args.figure_only:figure_only()
    elif args.build:build(args)
    elif args.finalize:finalize()
    else:p.error('Choose --figure-only, --build or --finalize')
if __name__=='__main__':main()
